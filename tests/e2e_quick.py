import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import h
path = h.setup("quick", TYPING_MAX_SECONDS="0.01")
from h import ok, run
import sqlite3

async def main():
    w = h.World(); B = w.B
    import db as D, quick as Q, config as C
    from form import parse_intro
    D.init_db(); D.add_moderator(50, "Иван", None, 1)
    assert C.FORM_VARIANT == "short"; ok("по умолчанию включён короткий диалог")
    con = lambda: sqlite3.connect(path)
    kb = lambda s: w.datas(s["kb"])
    shared = lambda uid, phone: w.feed(message={"message_id": 9000 + w.n, "date": 0, "chat": {"id": uid, "type": "private"}, "from": w.user(uid), "contact": {"phone_number": phone, "first_name": f"U{uid}", "user_id": uid}})

    # ---------- /start: представление и выбор способа связи ----------
    await w.msg(21, "/start")
    t = [s["text"] for s in w.to(21)]
    assert t == [Q.TEXTS["hello"], Q.TEXTS["ask_method"]], t; ok("/start: две реплики — представление и «как удобнее связаться»")
    assert "помощник" in t[0] and "контрактную службу" in t[0]; ok("бот представляется помощником, а не человеком")
    assert len(w.actions) >= 2, w.actions; ok("«печатает…» перед каждой репликой: %d раз" % len(w.actions))
    n1 = [x["text"] for x in w.to(1)]
    assert len(n1) == 1 and "Новый запуск бота" in n1[0] and "@u21" in n1[0]; ok("админу — одно уведомление о новом запуске бота")
    assert kb(w.last_to(21)) == ["qm:call", "qm:wa", "qm:tg", "qm:max"]; ok("кнопки способа связи: звонок, WhatsApp, Telegram, Max")
    assert not con().execute("SELECT 1 FROM applications").fetchone(); ok("пока контакта нет, заявки нет")

    # ---------- звонок: номер текстом, заявка создаётся сразу ----------
    w.clear(); await w.press(21, "qm:call")
    assert w.last_to(21)["text"] == Q.TEXTS["ask_call"] and "Поделиться номером" in str(w.last_to(21)["kb"]); ok("«Позвонить» → просьба о номере, кнопка «Поделиться номером»")
    assert "обработку персональных данных" in w.last_to(21)["text"]; ok("согласие на обработку данных — в просьбе о контакте")
    f = con().execute("SELECT started_form_at, contact_at, completed_at FROM funnel WHERE user_id=21").fetchone()
    assert f[0] and f[1] and not f[2]; ok("в воронке: начал отвечать и дошёл до контакта, заявка ещё не подана")
    w.clear(); await w.msg(21, "8888888")
    assert w.last_to(21)["text"] == Q.TEXTS["bad_contact"]; ok("короткий номер: «Не получилось распознать контакт…»")
    assert not con().execute("SELECT 1 FROM applications").fetchone(); ok("пока номер неверный, заявка не создаётся")
    w.clear(); await w.msg(21, "+7 (999) 123-45-67")
    admin_card = [s for s in w.to(1) if "Заявка #" in (s["text"] or "")]
    assert admin_card, w.to(1)
    c = admin_card[-1]["text"]
    for part in ("Заявка #1", "🆕 Новая", "U21", "+79991234567", "@u21", "Звонок"):
        assert part in c, (part, c)
    ok("сразу после контакта админу пришла карточка: имя и @username, телефон, способ связи «Звонок»")
    t = [s["text"] for s in w.to(21)]
    assert t[0] == Q.TEXTS["got_contact"] and t[1] == Q.TEXTS["ask_age"], t; ok("клиенту: «контакт получил» и вопрос про возраст")
    assert kb(w.last_to(21)) == [f"qa:{i}" for i in range(5)]; ok("возраст — кнопки-диапазоны")
    note = [s for s in w.to(50) if "Поступила новая заявка" in (s["text"] or "")][-1]
    for secret in ("+7999", "U21"):
        assert secret not in note["text"], secret
    ok("модератору — только «Поступила новая заявка», данные клиента скрыты")
    app = D.get_application(1)
    assert app["status"] == "new" and app["phone"] == "+79991234567" and app["contact_pref"] == "Звонок" and app["age"] is None, app; ok("в базе: статус «Новая», контакт и способ связи, остальное пока пусто")
    f = con().execute("SELECT completed_at FROM funnel WHERE user_id=21").fetchone()[0]; assert f; ok("в воронке: заявка подана")

    # ---------- уточнения: возраст, город, пол ----------
    w.clear(); await w.press(21, "qa:1")
    assert w.last_to(21)["text"] == Q.TEXTS["ask_city"]; ok("возраст кнопкой → вопрос про город")
    app = D.get_application(1); assert app["age_range"] == "25–34" and app["age"] is None; ok("в заявке диапазон возраста «25–34»")
    e = [x for x in w.edits if x["chat"] == 1][-1]["text"]; assert "25–34" in e; ok("карточка админа обновилась")
    w.clear(); await w.msg(21, "иркутск")
    assert w.last_to(21)["text"] == Q.TEXTS["ask_gender"] and kb(w.last_to(21)) == ["qg:0", "qg:1"]; ok("город текстом («иркутск» → «Иркутск») → вопрос про пол кнопками")
    w.clear(); await w.press(21, "qg:0")
    t = [s["text"] for s in w.to(21)]
    assert t == [Q.TEXTS["done"], Q.TEXTS["extra_invite"]], t; ok("клиенту: «Спасибо, заявка принята!…» + приглашение дополнить")
    assert "рассказать о себе подробнее" in t[1] and "опыт, образование, пожелания" in t[1]; ok("в приглашении — «рассказать о себе подробнее: опыт, образование, пожелания»")
    app = D.get_application(1)
    assert app["city"] == "Иркутск" and app["gender"] == "Мужской" and app["age_range"] == "25–34"; ok("в базе: город, пол, возраст")
    e = [x for x in w.edits if x["chat"] == 1][-1]["text"]
    assert "Иркутск" in e and "Мужской" in e and "25–34" in e and "Звонок" in e; ok("карточка админа показывает всё, включая способ связи")

    # ---------- дополнение информации обычным сообщением ----------
    w.clear(); await w.msg(21, "Служил срочную в ВДВ, хотел бы в десант")
    assert w.last_to(21)["text"] == Q.TEXTS["extra_ack"]; ok("любое сообщение после подачи: «Записал и передам специалисту»")
    assert D.get_application(1)["comment"] == "Служил срочную в ВДВ, хотел бы в десант"; ok("записано в комментарий клиента")
    await w.msg(21, "Есть мед. образование, фельдшер"); assert D.get_application(1)["comment"].count("\n") == 1; ok("несколько дополнений дописываются друг за другом")
    w.clear(); await w.msg(21, "ок"); assert "на рассмотрении" in w.last_to(21)["text"] and D.get_application(1)["comment"].count("\n") == 1; ok("слишком короткие реплики («ок») в заявку не пишутся")
    await w.msg(21, "/start"); assert "на рассмотрении" in w.last_to(21)["text"]; ok("повторный /start: «заявка на рассмотрении», новой не создать")
    await w.press(50, "mp:take"); mc = w.last_to(50)["text"]
    assert "Служил срочную" in mc and "фельдшер" in mc and "+79991234567" in mc and "Звонок" in mc, mc; ok("модератор, взяв заявку, видит контакт, способ связи и всё, что дописал клиент")
    t = __import__("table").row_block(D.get_application(1))
    assert "Клиент пишет:" in t and "Связь: Звонок" in t and "25–34 лет" in t; ok("в таблице: «Связь: Звонок», возраст диапазоном, «Клиент пишет: …»")
    from openpyxl import load_workbook
    import io
    w.docs.clear(); await w.press(50, "tb:xl"); ws = load_workbook(io.BytesIO(w.docs[-1][2])).active
    assert "Клиент пишет:" in [r for r in ws.iter_rows(values_only=True)][1][2]; ok("и в Excel «Моя таблица»")
    w.docs.clear(); await w.msg(1, "/export"); 
    if w.docs:
        rows = list(load_workbook(io.BytesIO(w.docs[-1][2])).active.iter_rows(values_only=True))
        assert "25–34" in rows[1]; ok("в админском Excel возраст диапазоном")

    # ---------- WhatsApp: иностранный номер, потом возраст числом ----------
    w.clear(); await w.msg(31, "/start"); w.clear(); await w.press(31, "qm:wa")
    assert w.last_to(31)["text"] == Q.TEXTS["ask_wa"] and "+380" in w.last_to(31)["text"]; ok("WhatsApp: просьба о номере с примером иностранного кода")
    await w.msg(31, "+380 50 123 45 67")
    app = [D.get_application(i) for i in (2,)][0]
    assert app["whatsapp"] == "+380501234567" and app["phone"] is None and app["contact_pref"] == "WhatsApp"; ok("номер +380… сохранён как WhatsApp без искажений")
    w.clear(); await w.msg(31, "17")
    assert "от 18 до 65" in w.last_to(31)["text"]; ok("возраст 17 числом: «доступна от 18 до 65 лет…»")
    await w.msg(31, "19"); assert w.last_to(31)["text"] == Q.TEXTS["ask_city"]; ok("возраст числом 19 принят → город")
    assert D.get_application(2)["age"] == 19; ok("точный возраст сохранён числом")
    await w.msg(31, "Самара, женщина")
    assert [s["text"] for s in w.to(31)][-2:] == [Q.TEXTS["done"], Q.TEXTS["extra_invite"]]; ok("город и пол одним сообщением — диалог завершён")
    app = D.get_application(2); assert app["city"] == "Самара" and app["gender"] == "Женский"; ok("город и пол сохранены")

    # ---------- Telegram: берётся @username, номер не нужен ----------
    w.clear(); await w.msg(41, "/start"); w.clear(); await w.press(41, "qm:tg")
    app = D.get_application(3)
    assert app and app["telegram"] == "@u41" and app["contact_pref"] == "Telegram" and app["phone"] is None; ok("Telegram: заявка создана по @username, номер не просили")
    assert Q.TEXTS["ask_age"] in [s["text"] for s in w.to(41)]; ok("и сразу — вопросы про возраст, город, пол")
    # номер прямо в ответ на «как удобнее» — звонок
    w.clear(); await w.msg(42, "/start"); w.clear(); await w.msg(42, "8 900 111-22-33")
    app = D.get_application(4); assert app["phone"] == "+79001112233" and app["contact_pref"] == "Звонок"; ok("номер вместо кнопки: считаем, что удобен звонок")
    # текст вместо выбора
    w.clear(); await w.msg(43, "/start"); w.clear(); await w.msg(43, "привет")
    assert w.last_to(43)["text"] == Q.TEXTS["pick_method"] and kb(w.last_to(43)) == ["qm:call", "qm:wa", "qm:tg", "qm:max"]; ok("слово вместо выбора: кнопки способа связи показаны снова")
    # Max с ником
    await w.press(43, "qm:max"); await w.msg(43, "ivan_max")
    app = D.get_application(5); assert app["max_contact"] == "ivan_max" and app["contact_pref"] == "Max"; ok("Max: ник сохраняется как есть")
    # контакт кнопкой
    w.clear(); await w.msg(44, "/start"); await w.press(44, "qm:call"); await shared(44, "79001112244")
    app = D.get_application(6); assert app["phone"] == "+79001112244"; ok("номер кнопкой «Поделиться номером»: заявка создана")

    # ---------- незавершённые уточнения ----------
    w.clear(); await w.msg(44, "Хочу в артиллерию, есть опыт срочной службы")
    assert w.last_to(44)["text"] == Q.TEXTS["ask_age"] or Q.TEXTS["unparsed"] in [s["text"] for s in w.to(44)]; ok("вместо ответа человек пишет о себе — бот повторяет вопрос")
    assert "артиллерию" in (D.get_application(6)["comment"] or ""); ok("и написанное попадает в заявку")
    w.clear(); await w.msg(44, "/cancel")
    assert w.last_to(44)["text"] == Q.TEXTS["skipped"]; ok("/cancel после создания заявки: вопросы пропускаются, заявка остаётся")
    assert D.get_application(6)["status"] == "new"; ok("заявка не отменена")

    # ---------- /start и /cancel до контакта ----------
    w.clear(); await w.msg(61, "/start"); w.clear(); await w.msg(61, "/start")
    assert [s["text"] for s in w.to(61)] == [Q.TEXTS["hello"], Q.TEXTS["ask_method"]]; ok("/start посреди диалога — начало заново")
    w.clear(); await w.msg(61, "/cancel"); assert "отменена" in w.last_to(61)["text"] and "apply" in str(w.datas(w.last_to(61)["kb"])); ok("/cancel до контакта: «Заявка отменена» и кнопка «Оставить заявку»")
    w.clear(); await w.press(61, "apply"); assert Q.TEXTS["hello"] in [s["text"] for s in w.to(61)]; ok("кнопка «Оставить заявку» (в т.ч. из напоминаний) запускает тот же короткий диалог")

    # ---------- разбор свободного текста ----------
    cases = [("27, Иркутск, мужчина", 27, "Иркутск", "Мужской"), ("Омск 34 ж", 34, "Омск", "Женский"),
             ("мне 41 год, живу в Казани, девушка", 41, "Казани", "Женский"), ("г. Санкт-Петербург, 30, м", 30, "Санкт-Петербург", "Мужской"),
             ("нижний новгород 28 муж", 28, "Нижний Новгород", "Мужской")]
    for text, age, city, gender in cases:
        assert parse_intro(text) == {"age": age, "city": city, "gender": gender}, (text, parse_intro(text))
    ok("разбор свободного текста работает по-прежнему")

    from utils import lead_priority
    assert lead_priority(None, None) == 1; ok("заявка без пояса и времени звонка: приоритет «любое время»")
    print("\nБЫСТРАЯ ЗАЯВКА — ВСЁ ПРОШЛО")
run(main())
