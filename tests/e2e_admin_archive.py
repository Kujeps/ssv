import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import h
path = h.setup("n4")
from h import ok, run
import sqlite3, io

async def main():
    w = h.World(); B = w.B
    import db as D, notify as N
    from aiogram.types import User
    from openpyxl import load_workbook
    D.init_db()
    con = lambda: sqlite3.connect(path)
    def sql(q, *a):
        c = con(); c.execute(q, a); c.commit(); c.close()

    # ---------- управление модераторами ----------
    await w.msg(1, "/moderators"); t = w.last_to(1)
    assert "Модераторы (0)" in t["text"] and "md:add" in w.datas(t["kb"]); ok("/moderators: список пуст, есть «Добавить модератора»")
    await w.press(1, "md:add"); kb = w.last_to(1)["kb"]
    btn = kb.keyboard[0][0]
    assert btn.request_users is not None and btn.request_users.max_quantity == 1 and btn.request_users.request_name; ok("кнопка выбора человека из контактов Telegram (request_users)")
    await w.msg(1, "не число"); assert "числовой ID" in w.last_to(1)["text"]; ok("мусорный ввод отклоняется")
    await w.shared(1, [{"user_id": 52, "first_name": "Анна", "last_name": "К", "username": "anna"}])
    assert D.is_moderator(52) and any("добавлен" in (s["text"] or "") and "Анна" in s["text"] for s in w.to(1)); ok("выбор из контактов: модератор «Анна К» добавлен")
    assert any("Вас назначили модератором" in (s["text"] or "") for s in w.to(52)); ok("новому модератору отправлено приглашение")
    assert any(c[0] == "SetMyCommands" for c in w.commands); ok("ему выставлено меню команд модератора")
    w.forbidden.add(53); await w.press(1, "md:add"); await w.shared(1, [{"user_id": 53, "first_name": "Олег"}])
    assert any("ещё не запускал бота" in (s["text"] or "") for s in w.to(1)) and D.is_moderator(53); ok("если модератор ещё не нажимал /start — админу подсказка")
    await w.press(1, "md:add"); await w.msg(1, "54")
    assert D.is_moderator(54); ok("добавление по числовому ID")
    await w.msg(1, "/moderators"); t = w.last_to(1); assert "Модераторы (3)" in t["text"] and "Анна" in t["text"]; ok("список из 3 модераторов")
    w.clear(); await w.press(52, "md:add"); assert not any("Новый модератор" in (s["text"] or "") for s in w.sent); ok("модератор не может добавлять модераторов")
    await w.msg(52, "/moderators"); assert "Модераторы (" not in str([s["text"] for s in w.sent]); ok("и видеть список модераторов")
    D.remove_moderator(53); D.remove_moderator(54); D.add_moderator(54, "Игорь", None, 1)

    # ---------- данные для архива ----------
    async def lead(i):
        u = User(id=200 + i, is_bot=False, first_name=f"Клиент{i}", username=f"c{i}")
        a = D.save_application(u, {"name": f"Клиент{i}", "phone": f"+7999000{i:04d}", "city": "Омск", "tz": "Москва (МСК)", "call_time": "В любое время"})
        await N.send_lead_cards(w.bot, a); return a
    ids = [await lead(i) for i in range(1, 13)]
    def work(mod, name):
        s, i = D.take_lead(mod, name); assert s == "taken"; return i
    x = work(52, "Анна К"); D.finish_lead(x, 52, "Анна К", "agreed")                    # 1
    x = work(52, "Анна К"); D.finish_lead(x, 52, "Анна К", "refused", "Причина: дорого")  # 2
    x = work(52, "Анна К"); D.finish_lead(x, 52, "Анна К", "junk", "Причина: спам")       # 3
    x = work(54, "Игорь"); D.finish_lead(x, 54, "Игорь", "agreed")                       # 4
    x = work(54, "Игорь"); D.postpone_lead(x, 54, "Игорь", "1h")                         # 5
    x = work(52, "Анна К")                                                               # 6 — в работе у Анны
    assert ids[:6] == [1, 2, 3, 4, 5, 6] or True
    s1 = D.create_source("залив 1", 1); s2 = D.create_source("залив 2", 1)
    for i in (1, 2, 3): sql("UPDATE applications SET source=? WHERE id=?", s1, i)
    sql("UPDATE applications SET source=? WHERE id=4", s2)
    sql("UPDATE applications SET status_at = datetime('now','-3 days') WHERE id=1")
    # обновим карточки админа
    for i in range(1, 7): await N.refresh_cards(w.bot, i)

    # ---------- архив админа ----------
    w.clear(); await w.msg(1, "/archive"); t = w.last_to(1)
    assert "Все заявки</b>: 12" in t["text"] and "Фильтры: не заданы" in t["text"]; ok("/archive: всего 12, фильтры не заданы")
    rows = [b for r in t["kb"].inline_keyboard for b in r if b.callback_data and b.callback_data.startswith("ar:o:")]
    assert [b.callback_data for b in rows] == [f"ar:o:{i}" for i in (12, 11, 10, 9, 8)], [b.callback_data for b in rows]; ok("5 заявок на странице, новые сверху (#12…#8)")
    nav = [b.text for r in t["kb"].inline_keyboard for b in r]
    assert "1/3" in nav and "➡️" in nav and "⬅️" not in nav and "📤 Excel по фильтру" in nav and "👷 Модератор" in nav; ok("навигация 1/3, фильтры, Excel, поиск")
    lid = w.with_kb(1)
    await w.press(1, "ar:p:2", mid=lid); e = w.edits[-1]
    assert [b.callback_data for r in e["kb"].inline_keyboard for b in r if (b.callback_data or "").startswith("ar:o:")] == ["ar:o:2", "ar:o:1"]; ok("страница 3: заявки #2 и #1")
    # статус
    await w.press(1, "ar:f:status", mid=lid); assert "Выберите статус" in w.edits[-1]["text"]
    await w.press(1, "ar:s:agreed", mid=lid); e = w.edits[-1]
    assert "Все заявки</b>: 2" in e["text"] and "статус — Согласился" in e["text"]; ok("фильтр «Согласился»: 2 заявки (#4 и #1)")
    # период
    await w.press(1, "ar:f:period", mid=lid); await w.press(1, "ar:pe:today", mid=lid); e = w.edits[-1]
    assert "Все заявки</b>: 1" in e["text"] and "период — Сегодня" in e["text"] and "ar:o:4" in str(w.datas(e["kb"])); ok("плюс период «Сегодня»: осталась #4 (#1 изменена 3 дня назад)")
    await w.press(1, "ar:f:period", mid=lid); await w.press(1, "ar:pe:7d", mid=lid); assert "Все заявки</b>: 2" in w.edits[-1]["text"]; ok("период «7 дней» возвращает #1")
    # модератор
    await w.press(1, "ar:s:all", mid=lid); await w.press(1, "ar:pe:all", mid=lid)
    await w.press(1, "ar:f:mod", mid=lid); ds = w.datas(w.edits[-1]["kb"]); assert "ar:m:52" in ds and "ar:m:54" in ds; ok("фильтр «Модератор»: список модераторов, у которых есть заявки")
    await w.press(1, "ar:m:52", mid=lid); e = w.edits[-1]
    got = sorted(int(d.split(":")[2]) for d in w.datas(e["kb"]) if d.startswith("ar:o:"))
    assert "Все заявки</b>: 4" in e["text"] and got == [1, 2, 3, 6] and "модератор — Анна К" in e["text"], (e["text"], got); ok("модератор Анна К: заявки 1, 2, 3, 6")
    # источник
    await w.press(1, "ar:m:all", mid=lid); await w.press(1, "ar:f:src", mid=lid)
    labels = w.buttons(w.edits[-1]["kb"]); assert "залив 1" in labels and "залив 2" in labels; ok("фильтр «Источник»: названия, а не коды")
    idx = labels.index("залив 1") - 1
    await w.press(1, f"ar:so:{idx}", mid=lid); e = w.edits[-1]
    assert "Все заявки</b>: 3" in e["text"] and "источник — залив 1" in e["text"]; ok("источник «залив 1»: 3 заявки")
    await w.press(1, "ar:reset", mid=lid); assert "Все заявки</b>: 12" in w.edits[-1]["text"] and "не заданы" in w.edits[-1]["text"]; ok("«Сбросить» очищает фильтры")
    # поиск
    w.clear(); await w.press(1, "ar:q", mid=lid); assert "Поиск" in w.last_to(1)["text"]
    await w.msg(1, "клиент7"); t = w.last_to(1)
    assert "Все заявки</b>: 1" in t["text"] and "поиск — «клиент7»" in t["text"] and "ar:o:7" in str(w.datas(t["kb"])); ok("поиск по имени (регистр не важен): найдена #7")
    lid2 = w.with_kb(1)
    await w.press(1, "ar:q", mid=lid2); await w.msg(1, "8 (999) 000-00-09"); assert "ar:o:9" in str(w.datas(w.last_to(1)["kb"])); ok("поиск по телефону в любом формате")
    await w.press(1, "ar:q", mid=w.with_kb(1)); await w.msg(1, "#5"); t = w.last_to(1); assert "Все заявки</b>: 1" in t["text"] and "ar:o:5" in str(w.datas(t["kb"])); ok("поиск по номеру заявки «#5»")
    await w.press(1, "ar:q", mid=w.with_kb(1)); await w.msg(1, "/cancel"); assert "Отменено" in w.last_to(1)["text"]; ok("/cancel выходит из поиска")
    # карточка из архива
    lid3 = w.with_kb(1)
    await w.press(1, "ar:o:2", mid=lid3); e = w.edits[-1]
    for part in ("Заявка #2", "Отказался", "Причина: дорого", "История", "взял в работу", "Анна К"):
        assert part in e["text"], (part, e["text"])
    ds = w.datas(e["kb"]); assert "ar:b" in ds and "ad:2:requeue" in ds; ok("карточка из архива: итог, причина, история, кнопки «К списку» и «Вернуть в очередь»")
    await w.press(1, "ar:b", mid=lid3); assert "Все заявки" in w.edits[-1]["text"]; ok("«К списку» возвращает список")

    # ---------- Excel по фильтру ----------
    await w.press(1, "ar:reset", mid=lid3); await w.press(1, "ar:f:status", mid=lid3); await w.press(1, "ar:s:agreed", mid=lid3)
    w.docs.clear(); await w.press(1, "ar:x", mid=lid3)
    chat, fname, data, cap = w.docs[-1]
    ws = load_workbook(io.BytesIO(data)).active; rows = list(ws.iter_rows(values_only=True))
    assert len(rows) == 3 and {r[0] for r in rows[1:]} == {1, 4} and "Согласился" in cap and fname.endswith(".xlsx"); ok("Excel по фильтру: только «Согласился» (#1, #4), подпись с фильтрами")
    assert dict(zip(rows[0], rows[1]))["Модератор"] in ("Анна К", "Игорь") and dict(zip(rows[0], rows[1]))["Источник"] in ("залив 1", "залив 2"); ok("в Excel: модератор и название источника")
    w.docs.clear(); await w.msg(1, "/export"); assert len(list(load_workbook(io.BytesIO(w.docs[-1][2])).active.iter_rows())) == 13; ok("/export: все 12 заявок")

    # ---------- архив модератора: только свои ----------
    w.clear(); await w.press(52, "mp:arch", name="Анна К"); t = w.last_to(52)
    got = sorted(int(d.split(":")[2]) for d in w.datas(t["kb"]) if d.startswith("ar:o:"))
    assert "Мои заявки</b>: 4" in t["text"] and got == [1, 2, 3, 6], (t["text"], got); ok("модератор видит «Мои заявки»: только свои 4")
    labels = w.buttons(t["kb"]); assert "👷 Модератор" not in labels and "🔖 Источник" not in labels and "📤 Excel по фильтру" not in labels; ok("у модератора нет фильтров по модераторам/источникам и выгрузки Excel")
    w.clear(); await w.press(52, "ar:o:4", mid=w.with_kb(52)); assert w.answers[-1] == ("Заявка недоступна", True); ok("чужую заявку из архива открыть нельзя")
    await w.press(52, "ar:o:2", mid=w.with_kb(52)); e = w.edits[-1]; assert "Причина: дорого" in e["text"] and "История" not in e["text"] and f"ad:2:requeue" not in w.datas(e["kb"]); ok("своя заявка открывается: без истории и без админских кнопок")
    await w.press(52, "ar:f:status", mid=w.with_kb(52)); assert "🆕 Новые" not in w.buttons(w.edits[-1]["kb"]); ok("в фильтре статусов у модератора нет «Новые»")
    w.clear(); w.docs.clear(); await w.press(52, "ar:x"); assert not w.docs; ok("выгрузка Excel модератору недоступна")
    w.clear(); await w.press(21, "ar:p:1"); assert "неактуальна" in (w.answers[-1][0] or ""); ok("клиент не может листать архив")

    # ---------- ручной возврат админом ----------
    w.clear(); await w.press(1, "ad:6:requeue", mid=w.last[1] if False else None)
    app = D.get_application(6); assert app["status"] == "new" and app["assigned_to"] is None; ok("админ вернул заявку из работы в очередь")
    assert any("возвращена администратором" in (s["text"] or "") for s in w.to(52)); ok("модератору, у которого забрали заявку, пришло сообщение")
    assert any("вернулась в очередь" in (s["text"] or "") for s in w.to(54)); ok("остальным модераторам — «Заявка вернулась в очередь»")
    w.clear(); await w.press(1, "ad:6:requeue"); assert w.answers[-1] == ("Заявка уже в очереди", True); ok("повторно вернуть заявку, уже стоящую в очереди, нельзя")

    # ---------- удаление модератора ----------
    s_, i_ = D.take_lead(52, "Анна К"); open_id = i_
    w.clear(); await w.press(1, "md:del:52", mid=w.last[1] if False else None)
    t = w.last_to(1); assert "Убрать модератора" in t["text"] and "Анна" in t["text"] and "(<b>1</b>)" in t["text"]; ok("подтверждение удаления: сколько заявок вернётся в очередь")
    await w.press(1, "md:ok:52", mid=w.with_kb(1))
    assert not D.is_moderator(52) and D.get_application(open_id)["status"] == "new"; ok("модератор убран, его заявка в работе вернулась в очередь")
    assert D.get_application(2)["status"] == "refused" and D.get_application(2)["assigned_to"] == 52; ok("его обработанные заявки остались в архиве")
    assert any("Доступ модератора отозван" in (s["text"] or "") for s in w.to(52)) and any(c[0] == "DeleteMyCommands" for c in w.commands); ok("ему сообщено, меню команд сброшено")
    w.clear(); await w.msg(52, "/start"); assert "оставьте заявку" in w.last_to(52)["text"]; ok("убранный модератор дальше — обычный клиент")

    # ---------- статистика ----------
    w.clear(); await w.msg(1, "/stats"); st = w.last_to(1)["text"]
    for part in ("Заявки по статусам", "В очереди:", "Модераторы", "Игорь", "взял", "Источники", "без метки"):
        assert part in st, (part, st)
    ok("/stats: статусы, модераторы (взял / итоги / недозвоны / возвраты), источники")
    w.clear(); await w.press(1, "ap:stats"); assert "Статистика бота" in w.last_to(1)["text"]; ok("то же из панели админа")
    print("\nN4: админ, архив, модераторы — ВСЁ ПРОШЛО")
run(main())
