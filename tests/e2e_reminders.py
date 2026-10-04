import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import h
path = h.setup("rem", FORM_VARIANT="full")
from h import ok, run
import io, sqlite3
from datetime import datetime, timedelta, timezone

async def main():
    w = h.World(); B = w.B
    import db as D, reminders as R, jobs as J, config as C
    from aiogram.types import User
    from openpyxl import load_workbook
    D.init_db(); D.add_moderator(50, "Иван", None, 1)
    con = lambda: sqlite3.connect(path)
    def sql(q, *a):
        c = con(); c.execute(q, a); c.commit(); c.close()
    def user(uid, started=None, form=False, **extra):
        u = User(id=uid, is_bot=False, first_name=f"Люд{uid}", username=f"p{uid}")
        D.touch_funnel(u, "started_at")
        if started is not None:
            sql("UPDATE funnel SET started_at = datetime('now', ?) WHERE user_id=?", f"-{started} hours", uid)
        if form:
            sql("UPDATE funnel SET started_form_at = datetime('now','-30 minutes') WHERE user_id=?", uid)
        for k, v in extra.items():
            sql(f"UPDATE funnel SET {k} = ? WHERE user_id=?", v, uid)
    def msgs(uid): return [s for s in w.sent if s["chat"] == uid]

    # --- график ---
    assert C.REMINDER_DELAYS_HOURS == [1, 5, 18, 24, 24, 24] and C.REMINDER_WINDOW_MSK == (10, 20)
    import staff as S; assert S.schedule_text() == "1, 6, 24, 48, 72 и 96 ч после запуска бота"; ok("график: 1 ч, 6 ч, 24 ч, затем ещё 3 раза через каждые 24 ч (1, 6, 24, 48, 72, 96)")
    assert len(R.TEXTS["start"]) == 6 and len(R.TEXTS["form"]) == 6
    t3 = R.TEXTS["start"][2]; assert "Остались вопросы?" in t3 and "ответим на все ваши вопросы" in t3; ok("текст суточного напоминания: «Остались вопросы? … ответим на все ваши вопросы»")
    assert not D.reminders_enabled(); ok("по умолчанию тумблер напоминаний ВЫКЛЮЧЕН")

    # --- /start больше не шлёт админу сообщений, но пишется в журнал ---
    await w.msg(21, "/start abc"); await w.msg(21, "/start")
    assert not w.to(1); ok("/start клиента: админу сообщений нет (ни первый, ни повторный)")
    rows = con().execute("SELECT user_id, source, is_new FROM start_log ORDER BY id").fetchall()
    assert rows == [(21, "abc", 1), (21, None, 0)]; ok("журнал: запуск с источником и повторный запуск")
    f = con().execute("SELECT starts, last_start_at FROM funnel WHERE user_id=21").fetchone(); assert f[0] == 2 and f[1]; ok("в воронке: число запусков и время последнего")

    # --- участники для проверки ---
    user(31, started=0.5)                         # A: запустил 30 мин назад — рано
    user(32, started=2)                           # B: 2 ч назад — пора, 1-е
    user(33, started=7, rem_count=1, rem_last_at=None)   # C: 7 ч назад, одно напоминание
    sql("UPDATE funnel SET rem_last_at = datetime('now','-6 hours') WHERE user_id=33")
    user(34, started=3, form=True)                # D: начал анкету и бросил — «form»
    user(35, started=3, rem_off=1)                # E: отписался
    user(36, started=3, blocked_at="2026-01-01 00:00:00")   # F: заблокировал бота
    user(37, started=3, completed_at="2026-01-01 00:00:00") # G: уже подал заявку
    user(1, started=5); user(50, started=5)       # админ и модератор
    user(38, started=240)                         # H: старый пользователь (10 суток назад)
    R_orig_window = R.REMINDER_WINDOW_MSK; R.REMINDER_WINDOW_MSK = (0, 24)

    # --- выключено: ничего не отправляется ---
    w.clear(); n = await R.run_reminders(w.bot); assert n == 0 and not w.sent; ok("при выключенном тумблере ничего не отправляется")

    # --- включение админом ---
    await w.msg(1, "/start"); lab = w.buttons(w.last_to(1)["kb"]); assert "🔔 Напоминания: ВЫКЛ" in lab and "👤 Пользователи" in lab; ok("в панели админа: «Напоминания: ВЫКЛ» и «Пользователи»")
    w.clear(); await w.press(1, "ap:rem"); t = w.last_to(1)
    assert "Включить напоминания?" in t["text"] and "1, 6, 24, 48, 72 и 96 ч" in t["text"] and "с 10:00 до 20:00 по Москве" in t["text"], t["text"]
    people, due = R.pending_counts(); assert (people, due) == (6, 4), (people, due)
    assert f"<b>{people}</b> чел., из них <b>{due}</b>" in t["text"]; ok(f"подтверждение: серия положена {people} чел., сейчас пора {due} (без админа, модератора, отписавшихся, заблокировавших и подавших)")
    assert D.reminders_enabled() is False; ok("пока не подтверждено — выключено")
    await w.press(1, "ap:remon", mid=w.with_kb(1)); assert D.reminders_enabled() and "ВКЛ" in str(w.buttons(w.last_to(1)["kb"])); ok("«Включить» → тумблер ВКЛ, панель обновилась")

    # --- первая рассылка ---
    w.clear(); n = await R.run_reminders(w.bot)
    got = {s["chat"]: s for s in w.sent}
    assert n == 4 and set(got) == {32, 33, 34, 38}, (n, set(got)); ok("отправлено 4: B, C, D и старому пользователю H; остальным — нет")
    assert got[32]["text"] == R.TEXTS["start"][0] and got[33]["text"] == R.TEXTS["start"][1] and got[38]["text"] == R.TEXTS["start"][0]; ok("тексты по порядку: B и H — 1-е, C — 2-е")
    assert got[34]["text"] == R.TEXTS["form"][0] and "Вы почти закончили" in got[34]["text"]; ok("бросившему анкету — свой текст «Вы почти закончили»")
    assert w.datas(got[32]["kb"]) == ["apply", "rem:off"] and "Не напоминать" in str(w.buttons(got[32]["kb"])); ok("под сообщением кнопки «Оставить заявку» и «Не напоминать»")
    assert w.buttons(got[34]["kb"])[0] == "📝 Заполнить заявку"
    for uid in (1, 50, 31, 35, 36, 37): assert uid not in got, uid
    ok("не получили: админ, модератор, «рано», отписавшийся, заблокировавший, подавший")
    w.clear(); assert await R.run_reminders(w.bot) == 0; ok("повторный запуск сразу же ничего не дублирует")

    # --- серия до конца ---
    texts = [R.TEXTS["start"][0]]
    for stage, gap in enumerate([5, 18, 24, 24, 24], start=1):
        sql("UPDATE funnel SET rem_last_at = datetime('now', ?) WHERE user_id=32", f"-{gap} hours")
        w.clear(); await R.run_reminders(w.bot)
        m = [s for s in w.sent if s["chat"] == 32]; assert len(m) == 1 and m[0]["text"] == R.TEXTS["start"][stage], (stage, m)
        texts.append(m[0]["text"])
    ok("серия: 6 разных сообщений — через 1, 5, 18, 24, 24, 24 ч после предыдущего")
    sql("UPDATE funnel SET rem_last_at = datetime('now','-100 hours') WHERE user_id=32")
    w.clear(); await R.run_reminders(w.bot); assert not [s for s in w.sent if s["chat"] == 32]; ok("после 6-го напоминания больше не пишем")
    assert len(set(texts)) == 6; ok("тексты не повторяются")

    # --- отписка ---
    mid = got[33]["id"]
    await w.press(33, "rem:off", mid=mid)
    assert w.answers[-1][0] == "Хорошо, больше не напомним" and "больше не будем напоминать" in w.edits[-1]["text"] and w.edits[-1]["kb"] is None; ok("«Не напоминать»: сообщение заменено подтверждением, кнопок нет")
    assert con().execute("SELECT rem_off FROM funnel WHERE user_id=33").fetchone()[0] == 1
    sql("UPDATE funnel SET rem_last_at = datetime('now','-30 hours') WHERE user_id=33")
    w.clear(); await R.run_reminders(w.bot); assert not [s for s in w.sent if s["chat"] == 33]; ok("отписавшемуся напоминаний больше нет")

    # --- заблокировал бота ---
    user(39, started=3); w.forbidden.add(39)
    w.clear(); n = await R.run_reminders(w.bot)
    assert con().execute("SELECT blocked_at FROM funnel WHERE user_id=39").fetchone()[0] and n == 0; ok("пользователь заблокировал бота: помечен и больше не получит")
    w.forbidden.clear()

    # --- подал заявку после напоминания: стоп и конверсия ---
    await w.press(34, "apply"); assert "Как к вам обращаться" in w.last_to(34)["text"]; ok("кнопка «Заполнить заявку» из напоминания запускает анкету")
    sql("UPDATE funnel SET completed_at = datetime('now') WHERE user_id=34")
    sql("UPDATE funnel SET rem_last_at = datetime('now','-30 hours') WHERE user_id=34")
    w.clear(); await R.run_reminders(w.bot); assert not [s for s in w.sent if s["chat"] == 34]; ok("после подачи заявки напоминания прекращаются")
    st = D.reminder_stats(); assert st["enabled"] and st["converted"] == 1 and st["optout"] >= 2 and st["total"] == 9, st; ok("статистика: отправлено %d, получили %d, подали после напоминания %d" % (st["total"], st["reminded"], st["converted"]))
    w.clear(); await w.msg(1, "/stats"); txt = w.last_to(1)["text"]
    for part in ("Напоминания</b> — включены", "Отправлено:", "1-е:", "Получили хотя бы одно", "после этого подали заявку: <b>1</b>", "Отписались"):
        assert part in txt, (part, txt)
    ok("/stats: блок «Напоминания» с конверсией")

    # --- тихие часы ---
    R.REMINDER_WINDOW_MSK = R_orig_window
    msk = timezone(timedelta(hours=3))
    for hh, mm, expect in ((9, 59, False), (10, 0, True), (19, 59, True), (20, 0, False), (3, 0, False)):
        now = datetime(2026, 10, 5, hh, mm, tzinfo=msk).astimezone(timezone.utc)
        assert R.in_window(now) is expect, (hh, mm)
    ok("окно отправки 10:00–20:00 МСК: 09:59 нет, 10:00 да, 19:59 да, 20:00 нет, 03:00 нет")
    user(40, started=5); w.clear()
    night = datetime(2026, 10, 5, 3, 0, tzinfo=msk).astimezone(timezone.utc)
    assert await R.run_reminders(w.bot, now=night + timedelta(days=400)) == 0 and not w.sent; ok("ночью (03:00 МСК) напоминания не уходят, даже если срок наступил")
    day = datetime(2026, 10, 5, 12, 0, tzinfo=msk).astimezone(timezone.utc) + timedelta(days=400)
    n = await R.run_reminders(w.bot, now=day); assert n >= 1 and [s for s in w.sent if s["chat"] == 40]; ok("в 12:00 МСК — уходят")

    # --- порциями ---
    R.REMINDER_WINDOW_MSK = (0, 24); R.REMINDER_BATCH = 3
    for uid in range(60, 68): user(uid, started=3)
    w.clear(); n1 = await R.run_reminders(w.bot); n2 = await R.run_reminders(w.bot)
    assert n1 == 3 and n2 == 3; ok("не волной: не больше 3 за одну проверку (лимит порции)")

    # --- панель выключается ---
    await w.press(1, "ap:rem"); assert not D.reminders_enabled() and "ВЫКЛ" in str(w.buttons(w.last_to(1)["kb"])); ok("повторное нажатие тумблера выключает напоминания сразу")
    R.REMINDER_BATCH = 100; w.clear(); assert await R.run_reminders(w.bot) == 0; ok("после выключения рассылка стоит")

    # --- раздел «Пользователи» ---
    w.clear(); await w.press(1, "ap:users"); t = w.last_to(1)
    assert "Пользователи бота" in t["text"] and "Запустили:" in t["text"] and "Люд" in t["text"]; ok("«Пользователи»: сводка и последние запуски")
    assert "us:not_applied:0" in w.datas(t["kb"]) and "us:x:all" in w.datas(t["kb"]); ok("фильтры и выгрузка в Excel")
    assert "Люд1)" not in t["text"] and "(@p50)" not in t["text"]; ok("админ и модератор в списке не показываются")
    lid = w.with_kb(1)
    await w.press(1, "us:applied:0", mid=lid); e = w.edits[-1]; assert "Подали заявку</b>: 2" in e["text"] and "✅" in e["text"]; ok("фильтр «Подали заявку»")
    await w.press(1, "us:optout:0", mid=lid); assert "Отписались</b>: 2" in w.edits[-1]["text"] or "Отписались</b>: 1" in w.edits[-1]["text"]; ok("фильтр «Отписались»")
    await w.press(1, "us:all:1", mid=lid); assert "стр. 2/" in w.edits[-1]["text"]; ok("листание страниц")
    w.docs.clear(); await w.press(1, "us:x:not_applied", mid=lid)
    ws = load_workbook(io.BytesIO(w.docs[-1][2])).active; rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:3] == ("User ID", "Имя", "Username") and "Напоминаний" in rows[0] and len(rows) > 5; ok("Excel «Пользователи»: %d строк, колонки запусков и напоминаний" % (len(rows) - 1))
    w.clear(); await w.press(21, "ap:users"); await w.press(21, "ap:rem"); assert not [s for s in w.sent if "Пользователи бота" in (s["text"] or "")] and not D.reminders_enabled(); ok("клиент не может открыть раздел и включить напоминания")
    print("\nНАПОМИНАНИЯ — ВСЁ ПРОШЛО")
run(main())
