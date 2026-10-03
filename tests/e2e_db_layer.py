import os, sys, sqlite3, tempfile
SP = tempfile.mkdtemp(prefix="svo_bot_test_")
os.environ.update(BOT_TOKEN="1:T", ADMIN_USER_IDS="1", DB_PATH=f"{SP}/t_db.db")
if os.path.exists(os.environ["DB_PATH"]): os.remove(os.environ["DB_PATH"])
# «старая» БД времён группы: 'work' без модератора должна уйти в очередь
c = sqlite3.connect(os.environ["DB_PATH"])
c.execute("CREATE TABLE applications (id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT DEFAULT CURRENT_TIMESTAMP, user_id INTEGER, username TEXT, full_name TEXT, name TEXT, phone TEXT, telegram TEXT, max_contact TEXT, whatsapp TEXT, status TEXT DEFAULT 'new')")
for st in ("new", "work", "nocall", "agreed"):
    c.execute("INSERT INTO applications (user_id, name, status) VALUES (?,?,?)", (900, f"old-{st}", st))
c.commit(); c.close()
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import db as D, utils as U, config as C
from datetime import datetime, timezone, timedelta
ok = lambda t: print("  PASS", t)

D.init_db(); D.init_db()   # дважды: миграция идемпотентна
st = dict(sqlite3.connect(os.environ["DB_PATH"]).execute("SELECT name, status FROM applications").fetchall())
assert st == {"old-new": "new", "old-work": "new", "old-nocall": "new", "old-agreed": "agreed"}, st; ok("миграция: все необработанные старые заявки — в очереди, закрытые остались")
assert D.queue_count() == 3

class Usr:  # заменитель пользователя
    def __init__(s, i): s.id=i; s.username=f"u{i}"; s.full_name=f"U{i}"
# приоритеты: в 15:00 МСК
t = datetime(2026, 10, 3, 12, 0, tzinfo=timezone.utc)      # 15:00 МСК
assert U.lead_priority(0, "День (12–17)", t) == 0; ok("МСК, окно «День» в 15:00 → приоритет 0")
assert U.lead_priority(0, "Утро (9–12)", t) == 2 and U.lead_priority(0, "В любое время", t) == 1
assert U.lead_priority(2, "Вечер (17–21)", t) == 0, "Екатеринбург 17:00"; ok("Екатеринбург (МСК+2): у клиента 17:00, окно «Вечер» → 0")
assert U.lead_priority(9, "День (12–17)", t) == 3, "Камчатка 00:00"; ok("Камчатка: у клиента полночь → в самый конец")
assert U.call_time_display("Вечер (17–21)", 2) == "Вечер (17–21) у клиента = 15:00–19:00 МСК" and U.call_time_display("День (12–17)", 0) == "День (12–17) МСК"; ok("окно переводится в МСК")


def mk(uid, call, tz): return D.save_application(Usr(uid), {"name": f"Лид{uid}", "phone": "+79990001122", "call_time": call, "tz": tz, "city": "Омск", "age": "30"})
a = mk(1001, "Утро (9–12)", "Москва (МСК)")            # окно закрыто → 2
b = mk(1002, "Вечер (17–21)", "Москва (МСК)")          # окно закрыто → 2
cc = mk(1003, "День (12–17)", "Москва (МСК)")          # сейчас удобно → 0
k = mk(1004, "Вечер (17–21)", "Камчатка (МСК+9)")      # у клиента ночь → 3
s1, got = D.take_lead(50, "Иван", now=t)
assert (s1, got) == ("taken", cc), (s1, got); ok(f"выдача: сначала та, кому удобно сейчас (#{cc}), а не самая старая")
assert D.take_lead(50, "Иван", now=t) == ("busy", cc); ok("пока заявка не закрыта, вторую не выдаёт")
s2, got2 = D.take_lead(51, "Пётр", now=t); assert got2 == 1, got2; ok("второй модератор: дальше «любое время» (#1, старая без окна), не ночные")
assert not D.finish_lead(cc, 51, "Пётр", "agreed"); ok("чужую заявку закрыть нельзя")
assert D.finish_lead(cc, 50, "Иван", "refused", "не хочет"); ok("закрытие своей заявки с причиной")
assert D.get_application(cc)["notes"][0]["text"] == "не хочет"
# ночью — в самом конце
order = []
for mod in (60, 61, 62, 63, 64, 65):
    s_, i_ = D.take_lead(mod, f"M{mod}", now=t)
    order.append(i_)
assert order[-1] == k or k in order[-2:], order; assert D.queue_count() >= 0
print("   порядок выдачи:", order, "(ночной #%d — позже обычных)" % k)
assert order.index(k) > order.index(a) and order.index(k) > order.index(b); ok("заявка с ночью у клиента выдаётся после остальных")

# недозвон / перезвон
n = D.postpone_lead(1, 51, "Пётр", "1h"); assert n == 1; ok("недозвон: попытка 1, заявка осталась за модератором")
assert [p["id"] for p in D.parked_leads(51)] == [1] and D.active_lead(51) is None
assert D.take_lead(51, "Пётр", now=t)[0] in ("taken", "empty")
# перезвон
app = D.get_application(1); assert app["status"] == "nocall" and app["assigned_to"] == 51 and app["callback_at"]
for mod in (60, 61, 62, 63, 64, 65): pass
st_before = D.take_callback(51, "Пётр", 1)
print("   take_callback →", st_before)
# возврат с причиной
active51 = D.active_lead(51)
if active51:
    assert D.release_lead(active51["id"], 51, "Пётр", "не успел"); ok("возврат в очередь с причиной → снова в очереди")
    assert D.get_application(active51["id"])["status"] == "new" and D.get_application(active51["id"])["assigned_to"] is None
# таймеры
tt = datetime.now(timezone.utc)
ids = [r[0] for r in sqlite3.connect(os.environ["DB_PATH"]).execute("SELECT id FROM applications WHERE status='work'")]
assert ids, "должна быть хотя бы одна заявка в работе"
holder = D.get_application(ids[0])["assigned_to"]
assert D.expired_holds(tt) == []; ok("до истечения часа автовозврата нет")
later = tt + timedelta(minutes=61)
exp = D.expired_holds(later); assert any(i == ids[0] for i, _ in exp); ok("через 61 минуту заявка попадает под автовозврат")
warn = D.hold_warnings(tt + timedelta(minutes=52)); assert any(i == ids[0] and left <= 8 for i, _, left in warn); ok("через 52 мин — предупреждение (осталось ≤8)")
assert D.hold_warnings(tt + timedelta(minutes=30)) == []; ok("в середине срока предупреждений нет")
assert D.system_requeue(ids[0], "timeout", holder); assert D.get_application(ids[0])["status"] == "new"; ok("автовозврат: заявка снова в очереди")
# перезвон просрочен
hold2 = [r for r in D.parked_leads(51)]
if hold2:
    assert D.due_callbacks(tt) == [] and (hold2[0]["id"], 51) in D.due_callbacks(tt + timedelta(hours=2)); ok("пора перезвонить через 1 час")
    assert D.overdue_callbacks(tt + timedelta(hours=5)) == [] and D.overdue_callbacks(tt + timedelta(hours=26)); ok("просрочка 24 ч → возврат в очередь")
# удалить модератора
D.add_moderator(70, "Анна", "anna", 1); assert D.is_moderator(70); ok("модератор добавлен")
sx, ix = D.take_lead(70, "Анна", now=t); assert sx == "taken"
back = D.remove_moderator(70); assert back == [ix] and not D.is_moderator(70) and D.get_application(ix)["status"] == "new"; ok("удаление модератора возвращает его заявки в очередь")
# архив / поиск
tot, pg, rows = D.archive_page({}, 0); assert tot == 8 + 0 or tot > 0
tot_ref, _, rr = D.archive_page({"status": "refused"}, 0); assert [r["id"] for r in rr] == [cc]; ok("архив: фильтр по статусу")
tot, _, rr = D.archive_page({"q": "ЛИД1003"}, 0); assert [r["id"] for r in rr] == [cc]; ok("поиск по имени без учёта регистра (кириллица)")
tot, _, rr = D.archive_page({"q": "8 (999) 000-11-22"}, 0); assert tot >= 4; ok("поиск по телефону в любом формате")
tot, _, rr = D.archive_page({"q": f"#{cc}"}, 0); assert [r["id"] for r in rr] == [cc]; ok("поиск по номеру заявки (#N)")
tot, _, rr = D.archive_page({"mod": 50}, 0); assert [r["id"] for r in rr] == [cc]; ok("фильтр по модератору")
tot, _, rr = D.archive_page({}, 0, scope_mod=50); assert [r["id"] for r in rr] == [cc]; ok("модератор видит только свои заявки (scope)")
tot, _, rr = D.archive_page({"period": "today"}, 0); assert tot >= 4; ok("период «сегодня»")
tot, pg, rr = D.archive_page({}, 99, per_page=5); assert pg == (tot - 1) // 5; ok("страница за пределами → последняя")
# статистика и Excel
ms = {m["user_id"]: m for m in D.moderator_stats()}
assert ms[50]["refused"] == 1 and ms[50]["takes"] >= 1 and ms[51]["nocalls"] == 1; ok("статистика модераторов: возьмём/отказ/недозвон")
assert ms[51]["releases"] >= 0 and any(m["timeouts"] for m in ms.values()); ok("автовозвраты записаны на модератора")
from openpyxl import load_workbook
ws = load_workbook(D.build_export_xlsx({"status": "refused"})).active
rows = list(ws.iter_rows(values_only=True)); assert len(rows) == 2 and dict(zip(rows[0], rows[1]))["Модератор"] == "Иван"; ok("Excel по фильтру: только отказы, колонка «Модератор»")
print("\nБД: ВСЁ ПРОШЛО")
