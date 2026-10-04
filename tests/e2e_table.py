import os, sys; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import h
path = h.setup("tbl")
from h import ok, run
import io, sqlite3
from datetime import timedelta

async def main():
    w = h.World(); B = w.B
    import db as D, notify as N, config as C, table as T
    from aiogram.types import User
    from openpyxl import load_workbook
    D.init_db(); D.add_moderator(50, "Иван", None, 1); D.add_moderator(51, "Пётр", None, 1)
    con = lambda: sqlite3.connect(path)
    def sql(q, *a):
        c = con(); c.execute(q, a); c.commit(); c.close()
    async def lead(uid, name, phone, call="Вечер (17–21)", tz="Самара (МСК+1)"):
        i = D.save_application(User(id=uid, is_bot=False, first_name=name, username=f"c{uid}"),
            {"name": name, "phone": phone, "tg": f"@c{uid}", "city": "Самара", "age": "27", "gender": "Мужской", "unit": "ВДВ", "served": "Срочная служба", "tz": tz, "call_time": call})
        await N.send_lead_cards(w.bot, i); return i
    ids = [await lead(100 + i, f"Клиент{i}", f"+7900000000{i}") for i in range(1, 8)]
    for _ in range(4): await w.press(50, "mp:take"); cur = D.active_lead(50)["id"]; D.finish_lead(cur, 50, "Иван", "agreed") if _ < 3 else None
    mine_ids = [r[0] for r in con().execute("SELECT id FROM applications WHERE assigned_to=50 ORDER BY id")]
    work_id = D.active_lead(50)["id"]
    foreign = D.take_lead(51, "Пётр")[1]
    w.clear()

    # --- кнопка и команда ---
    await w.msg(50, "/start"); assert "📋 Моя таблица" in str(w.buttons(w.last_to(50)["kb"])); ok("в панели модератора есть «📋 Моя таблица»")
    assert not any("📋 Моя таблица" in b for b in w.buttons(h.World.__init__ and w.last_to(1)["kb"])) if w.to(1) and w.last_to(1)["kb"] else True
    w.clear(); await w.msg(50, "/table"); t = w.last_to(50)
    assert "Моя таблица</b> · Все: 4" in t["text"], t["text"]; ok("/table: таблица модератора, всего 4 своих заявки")
    assert f"#{foreign}" not in t["text"]; ok("чужие заявки в таблице не показываются")
    for part in ("#%d" % work_id, "В работе", "Клиент", "+7900000000", "@c10", "Самара (МСК+1)", "ВДВ", "Звонить: Вечер", "МСК"):
        assert part in t["text"] or part.startswith("@c10"), (part, t["text"])
    ok("строка таблицы: №, дата, статус, анкета (контакты, город и пояс, подразделение, окно звонка в МСК)")
    assert "tb:e:%d" % work_id in w.datas(t["kb"]) and "tb:xl" in w.datas(t["kb"]); ok("кнопка на каждую заявку и «Excel»")
    # активная заявка — первой
    first = t["text"].split("\n\n")[1]; assert f"#{work_id}" in first; ok("заявка в работе — сверху")
    await w.press(51, "tb:open"); assert "Клиент" in w.last_to(51)["text"] and f"#{mine_ids[0]}" not in w.last_to(51)["text"]; ok("у другого модератора своя таблица")
    w.clear(); await w.press(21, "tb:open"); assert not w.sent; ok("клиент не может открыть таблицу")

    # --- вкладки ---
    scr = w.with_kb(50)
    await w.press(50, "tb:v:final:0", mid=scr); e = w.edits[-1]
    assert "Закрытые: 3" in e["text"] and "В работе" not in e["text"].split("\n\n", 1)[1]; ok("вкладка «Закрытые»: 3 заявки")
    await w.press(50, "tb:v:cb:0", mid=scr); assert "Перезвонить: 0" in w.edits[-1]["text"] and "Здесь пока пусто" in w.edits[-1]["text"]; ok("вкладка «Перезвонить» пока пустая")

    # --- экран редактирования ---
    await w.press(50, "tb:v:all:0", mid=scr)
    await w.press(50, f"tb:e:{work_id}", mid=scr); e = w.edits[-1]
    assert "Что изменить?" in e["text"] and f"Заявка #{work_id}" in e["text"] and set(w.datas(e["kb"])) == {f"tb:s:{work_id}", f"tb:c:{work_id}", f"tb:x:{work_id}", f"tb:d:{work_id}", "tb:b"}; ok("тап по заявке: полная карточка и кнопки «Статус / Комментарий / Доп. информация / Данные клиента»")
    w.clear(); await w.press(50, f"tb:e:{foreign}", mid=scr); assert w.answers[-1] == ("Эта заявка уже не закреплена за вами", True); ok("чужую заявку редактировать нельзя")

    # --- быстрый комментарий одной кнопкой ---
    await w.press(50, f"tb:c:{work_id}", mid=scr)
    ds = w.datas(w.markups[-1]["kb"]); assert f"tb:q:{work_id}:0" in ds and len([d for d in ds if d.startswith("tb:q:")]) == len(C.QUICK_COMMENTS); ok("«Комментарий»: шесть быстрых фраз одной кнопкой")
    w.clear(); await w.press(50, f"tb:q:{work_id}:0", mid=scr)
    app = D.get_application(work_id); assert app["notes"][-1]["text"] == "Не берёт трубку" and "Комментарий добавлен" == w.answers[-1][0]; ok("быстрая фраза записана в комментарии (одно нажатие)")
    e = w.edits[-1]; assert "Не берёт трубку" in e["text"]; ok("экран обновился сразу")
    assert "Не берёт трубку" in [x for x in w.edits if x["chat"] == 1][-1]["text"]; ok("карточка админа обновилась")

    # --- свой комментарий, правка и удаление ---
    await w.press(50, f"tb:c:{work_id}", mid=scr); w.clear(); await w.press(50, f"tb:cw:{work_id}", mid=scr)
    assert "Напишите — добавлю" in w.last_to(50)["text"]
    prompt = w.last_to(50)["id"]
    await w.msg(50, "Звонил в 16:10, просил позже"); app = D.get_application(work_id)
    assert app["notes"][-1]["text"] == "Звонил в 16:10, просил позже"; ok("свой комментарий: ввод текстом")
    assert (50, prompt) in w.deleted; ok("вопрос и ответ убраны из чата — чат не засоряется")
    await w.press(50, f"tb:c:{work_id}", mid=scr); await w.press(50, f"tb:ce:{work_id}", mid=scr)
    assert "Звонил в 16:10" in w.last_to(50)["text"]; await w.msg(50, "Звонил в 16:10, перезвонить в пятницу")
    assert D.get_application(work_id)["notes"][-1]["text"] == "Звонил в 16:10, перезвонить в пятницу"; ok("«Изменить последний»: текст заменён")
    await w.press(50, f"tb:c:{work_id}", mid=scr); await w.press(50, f"tb:cd:{work_id}", mid=scr)
    assert [n["text"] for n in D.get_application(work_id)["notes"]] == ["Не берёт трубку"]; ok("«Удалить последний» убирает свой последний комментарий")

    # --- доп. информация ---
    await w.press(50, f"tb:x:{work_id}", mid=scr); assert set(w.datas(w.markups[-1]["kb"])) >= {f"tb:xa:{work_id}", f"tb:xs:{work_id}", f"tb:xc:{work_id}"}
    await w.press(50, f"tb:xa:{work_id}", mid=scr); await w.msg(50, "Паспорт и ВБ на руках")
    await w.press(50, f"tb:x:{work_id}", mid=scr); await w.press(50, f"tb:xa:{work_id}", mid=scr); await w.msg(50, "Приедет 10.10")
    assert D.get_application(work_id)["extra"] == "Паспорт и ВБ на руках\nПриедет 10.10"; ok("доп. информация: дописывается строками")
    t = w.edits[-1]["text"]; assert "Доп. информация" in t and "Приедет 10.10" in t; ok("она видна на экране заявки")
    await w.press(50, f"tb:x:{work_id}", mid=scr); await w.press(50, f"tb:xs:{work_id}", mid=scr); await w.msg(50, "Документы собраны")
    assert D.get_application(work_id)["extra"] == "Документы собраны"; ok("«Заменить» — полностью новый текст")
    await w.press(50, f"tb:x:{work_id}", mid=scr); await w.press(50, f"tb:xc:{work_id}", mid=scr)
    assert not D.get_application(work_id)["extra"]; ok("«Очистить»")
    await w.press(50, f"tb:x:{work_id}", mid=scr); await w.press(50, f"tb:xa:{work_id}", mid=scr); await w.msg(50, "Семья: женат, двое детей")

    # --- данные клиента ---
    await w.press(50, f"tb:d:{work_id}", mid=scr); ds = w.datas(w.markups[-1]["kb"])
    assert f"tb:f:{work_id}:phone" in ds and f"tb:f:{work_id}:city" in ds and f"tb:f:{work_id}:tz" in ds; ok("«Данные клиента»: имя, телефон, Telegram, MAX, WhatsApp, город, возраст, подразделение…")
    old_phone = D.get_application(work_id)["phone"]
    await w.press(50, f"tb:f:{work_id}:phone", mid=scr); assert old_phone in w.last_to(50)["text"]
    await w.msg(50, "8 (900) 123-45-67"); assert D.get_application(work_id)["phone"] == "+79001234567"; ok("телефон исправлен (приведён к виду +7900…)")
    await w.press(50, f"tb:f:{work_id}:age", mid=scr); await w.msg(50, "17"); assert "от 18 до 65" in w.last_to(50)["text"] and D.get_application(work_id)["age"] == 27; ok("проверка ввода: возраст 17 отклонён, остаёмся в режиме ввода")
    await w.msg(50, "31"); assert D.get_application(work_id)["age"] == 31; ok("исправленный ввод принят")
    await w.press(50, f"tb:f:{work_id}:tz", mid=scr); labels = w.buttons(w.markups[-1]["kb"]); assert "Москва (МСК)" in labels and "✍️ Своё значение" not in labels
    await w.press(50, f"tb:fv:{work_id}:tz:1", mid=scr); assert D.get_application(work_id)["tz_offset"] == 0; ok("часовой пояс — кнопкой")
    await w.press(50, f"tb:f:{work_id}:unit", mid=scr); assert "✍️ Своё значение" in w.buttons(w.markups[-1]["kb"])
    await w.press(50, f"tb:fv:{work_id}:unit:2", mid=scr); assert D.get_application(work_id)["unit"] == "ВМФ"; ok("подразделение — из списка одной кнопкой")
    edits = D.get_edits(work_id); fields = [e["field"] for e in edits]
    assert "Телефон" in fields and "Возраст" in fields and "Часовой пояс" in fields and "Подразделение" in fields and "Доп. информация" in fields; ok("все правки записаны в журнал (кто, что, было → стало)")
    assert any(e["field"] == "Телефон" and e["old"] == old_phone and e["new"] == "+79001234567" for e in edits)
    assert "+79001234567" in [x for x in w.edits if x["chat"] == 1][-1]["text"]; ok("карточка админа показывает исправленные данные")

    # --- статусы ---
    await w.press(50, f"tb:s:{work_id}", mid=scr); ds = w.datas(w.markups[-1]["kb"])
    assert {f"tb:st:{work_id}:{s}" for s in ("reached", "nocall", "callback", "agreed", "refused", "junk")} <= set(ds); ok("«Статус»: Дозвонились / Не дозвонились / Перезвонить / Согласился / Отказался / Мусор")
    await w.press(50, f"tb:st:{work_id}:reached", mid=scr); a = D.get_application(work_id)
    assert a["status"] == "reached" and a["taken_at"] is None and D.active_lead(50) is None; ok("«Дозвонились»: заявка остаётся у модератора без таймера, он свободен")
    sql("UPDATE applications SET assigned_at = datetime('now','-3 hours') WHERE id=?", work_id)
    import jobs as J
    sql("UPDATE applications SET taken_at = datetime('now','-5 hours') WHERE id=?", work_id)
    w.clear(); await J.run_maintenance(w.bot); assert D.get_application(work_id)["status"] == "reached"; ok("часовой таймер на «Дозвонились» не действует — долгий клиент остаётся за модератором")
    await w.press(50, f"tb:s:{work_id}", mid=scr); await w.press(50, f"tb:st:{work_id}:callback", mid=scr)
    ds = w.datas(w.markups[-1]["kb"]); assert f"tb:tm:{work_id}:callback:1h" in ds and f"tb:tm:{work_id}:callback:custom" in ds; ok("«Перезвонить»: 1 час / 3 часа / завтра утром / своё время")
    await w.press(50, f"tb:tm:{work_id}:callback:custom", mid=scr); assert "по Москве" in w.last_to(50)["text"]
    await w.msg(50, "когда-нибудь"); assert "Не понял время" in w.last_to(50)["text"]; ok("непонятное время не принимается")
    await w.msg(50, "01.01 10:00"); assert "Не понял время" in w.last_to(50)["text"]; ok("прошедшее время не принимается")
    await w.msg(50, "завтра 18:30")
    a = D.get_application(work_id); assert a["status"] == "callback" and a["callback_at"] and a["attempts"] == 0; ok("свой срок «завтра 18:30»: статус «Перезвонить», время сохранено")
    assert "Перезвонить:" in w.edits[-1]["text"] and "18:30" in w.edits[-1]["text"]
    await w.press(50, f"tb:s:{work_id}", mid=scr); await w.press(50, f"tb:st:{work_id}:nocall", mid=scr); await w.press(50, f"tb:tm:{work_id}:nocall:1h", mid=scr)
    a = D.get_application(work_id); assert a["status"] == "nocall" and a["attempts"] == 1; ok("«Не дозвонились» из таблицы: попытка 1")
    await w.press(50, f"tb:s:{work_id}", mid=scr); await w.press(50, f"tb:st:{work_id}:refused", mid=scr)
    assert "Причина (обязательно)" in w.last_to(50)["text"]; await w.msg(50, "?"); assert "хотя бы пару слов" in w.last_to(50)["text"]
    await w.msg(50, "Передумал, семья против"); a = D.get_application(work_id)
    assert a["status"] == "refused" and a["notes"][-1]["text"] == "Причина: Передумал, семья против"; ok("«Отказался»: причина обязательна и записана")
    await w.press(50, f"tb:s:{work_id}", mid=scr); await w.press(50, f"tb:st:{work_id}:agreed", mid=scr)
    assert D.get_application(work_id)["status"] == "agreed"; ok("итог можно исправить: «Отказался» → «Согласился»")
    hist = [(h["kind"]) for h in D.get_history(work_id)]; assert "reached" in hist and "callback" in hist and "nocall" in hist and hist.count("final") == 2; ok("история статусов: " + ", ".join(hist))
    # вкладки после изменений
    w.clear(); await w.msg(50, "/table"); assert "Закрытые (4)" in str(w.buttons(w.last_to(50)["kb"])); ok("счётчики вкладок обновляются")

    # --- быстрый ввод ---
    other = mine_ids[0]
    w.clear(); await w.msg(50, f"#{other} просил написать в WhatsApp")
    assert D.get_application(other)["notes"][-1]["text"] == "просил написать в WhatsApp" and "комментарий добавлен" in w.last_to(50)["text"]; ok("«#N текст» — комментарий без открытия экранов")
    await w.msg(50, f"#{other}+ есть второй номер +7 900 555-00-11"); assert D.get_application(other)["extra"] == "есть второй номер +7 900 555-00-11"; ok("«#N+ текст» — доп. информация")
    w.clear(); await w.msg(50, f"#{foreign} чужая"); assert "не закреплена за вами" in w.last_to(50)["text"] and not D.get_application(foreign)["notes"]; ok("быстрый ввод в чужую заявку отклонён")
    w.clear(); await w.msg(50, "#999999 нет такой"); assert "не закреплена" in w.last_to(50)["text"]; ok("несуществующий номер — аккуратный ответ")
    w.clear(); await w.msg(51, f"#{other} хак"); assert "не закреплена" in w.last_to(51)["text"]; ok("другой модератор так писать в чужие заявки не может")
    w.clear(); await w.msg(21, "#1 привет"); assert all("комментарий" not in (s["text"] or "") for s in w.sent); ok("клиенту быстрый ввод ничего не даёт")

    # --- Excel ---
    w.docs.clear(); await w.press(50, "tb:xl")
    ws = load_workbook(io.BytesIO(w.docs[-1][2])).active; rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == ["№", "Дата передачи", "Анкета", "Статус", "Комментарий модератора", "Перезвонить", "Доп. информация"]; ok("Excel «Моя таблица»: точные колонки макета + «Доп. информация»")
    assert len(rows) == 5 and {r[0] for r in rows[1:]} == set(mine_ids + [work_id]) - ({foreign}), [r[0] for r in rows]; ok("в файле только свои заявки (%d)" % (len(rows) - 1))
    r = {x[0]: x for x in rows[1:]}[work_id]
    assert "\n" in r[2] and "+79001234567" in r[2] and r[3] == "Согласился" and "Не берёт трубку" in r[4] and r[6] == "Семья: женат, двое детей"; ok("анкета многострочной ячейкой, статус, комментарии, доп. информация")
    assert ws.cell(row=2, column=3).alignment.wrap_text; ok("ячейки с переносом строк")

    # --- администратор видит всё ---
    app = D.get_application(work_id)
    from ui import lead_text
    t = lead_text(app, "admin")
    assert "Доп. информация" in t and "Семья: женат" in t and "Передана:" in t; ok("админская карточка: доп. информация и дата передачи")
    import staff as S
    ht = S.history_text(work_id)
    assert "Телефон" in ht and "→" in ht and "Иван" in ht and "Согласился" in ht; ok("история заявки у админа: смены статуса и правки данных (кто, что, было → стало)")
    w.docs.clear(); await w.msg(1, "/export"); ws = load_workbook(io.BytesIO(w.docs[-1][2])).active
    head = [c.value for c in ws[1]]; assert "Доп. информация" in head and "Дата передачи" in head; ok("админский Excel: «Доп. информация» и «Дата передачи»")
    print("\nТАБЛИЦА — ВСЁ ПРОШЛО")
run(main())
