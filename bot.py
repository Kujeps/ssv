import asyncio
import logging
import os
import re
import secrets
import sqlite3
import string
from contextlib import closing, suppress
from datetime import datetime, timedelta, timezone
from html import escape
from io import BytesIO

from aiogram import Bot, Dispatcher, F, Router
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ChatType, ParseMode
from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError, TelegramRetryAfter
from aiogram.filters import Command, CommandObject, CommandStart, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import (
    BotCommand,
    BotCommandScopeChat,
    BotCommandScopeDefault,
    BufferedInputFile,
    CallbackQuery,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    KeyboardButton,
    Message,
    ReactionTypeEmoji,
    ReplyKeyboardMarkup,
    ReplyKeyboardRemove,
    ReplyParameters,
    User,
)
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font

load_dotenv()


def env_ids(name: str) -> list[int]:
    return [int(x) for x in os.getenv(name, "").replace(" ", "").split(",") if x]


BOT_TOKEN = os.environ["BOT_TOKEN"]
# Чаты, куда уходят готовые заявки (группа менеджеров).
ADMIN_IDS = env_ids("ADMIN_CHAT_IDS")
ADMIN_CHAT_SET = set(ADMIN_IDS)
# Админы бота: /stats, /export, /broadcast и уведомления о /start (в личке с ботом).
ADMIN_USER_IDS = env_ids("ADMIN_USER_IDS")
# Кто может менять статусы заявок и оставлять заметки. Пусто = любой участник группы заявок.
MANAGER_IDS = env_ids("MANAGER_IDS")
DB_PATH = os.getenv("DB_PATH", "leads.db")
# Часовой пояс для времени в карточках и Excel (смещение от UTC, по умолчанию Москва).
LOCAL_TZ = timezone(timedelta(hours=int(os.getenv("TZ_OFFSET_HOURS", "3"))))
# Напоминать о необработанной заявке через N минут (0 — выключить), не более REMIND_MAX раз.
REMIND_AFTER_MIN = int(os.getenv("REMIND_AFTER_MIN", "120"))
REMIND_MAX = 3


# ---------- настройки анкеты (при необходимости правятся здесь) ----------

MIN_AGE, MAX_AGE = 18, 65
SOURCE_CODE_LEN = 24  # длина случайного кода в рекламной ссылке

UNITS = [
    "Сухопутные войска",
    "ВДВ",
    "ВМФ",
    "ВКС",
    "Артиллерия / ракетные",
    "Беспилотные системы",
    "Связь / РЭБ",
    "Медицинская служба",
    "Не определился — подскажите",
]
GENDERS = ["Мужской", "Женский"]
SERVED = ["Не служил(а)", "Срочная служба", "Служба по контракту"]
CALL_TIMES = ["Утро (9–12)", "День (12–17)", "Вечер (17–21)", "В любое время"]
MEDICAL = ["Да", "Нет"]
CHOICES = {
    "gender": GENDERS, "medical": MEDICAL, "unit": UNITS,
    "served": SERVED, "call_time": CALL_TIMES,
}

# Шаги, которые нельзя пропустить. Остальные можно (но нужен хотя бы один контакт).
REQUIRED_STEPS = {"name", "gender", "medical", "age", "city", "unit"}

# Пока заявка человека в одном из этих статусов, новую он подать не может.
# После «Согласился» / «Отказался» / «Мусор» повторная заявка снова разрешена.
BLOCKING_STATUSES = ("new", "work", "nocall")

# Статусы заявок: ключ -> (значок, название).
STATUSES = {
    "new": ("🆕", "Новая"),
    "work": ("🔧", "В работе"),
    "nocall": ("📵", "Не дозвонились"),
    "agreed": ("✅", "Согласился"),
    "refused": ("❌", "Отказался"),
    "junk": ("🗑", "Мусор"),
}
STATUS_BUTTONS = {
    "work": "🔧 Беру в работу",
    "nocall": "📵 Не дозвонились",
    "agreed": "✅ Согласился",
    "refused": "❌ Отказался",
    "junk": "🗑 Мусор / спам",
    "new": "↩️ Вернуть в новые",
}
STATUS_BUTTON_ORDER = ["work", "nocall", "agreed", "refused", "junk", "new"]


class Form(StatesGroup):
    name = State()
    phone = State()
    tg = State()
    max = State()
    wa = State()
    gender = State()
    medical = State()
    age = State()
    city = State()
    unit = State()
    served = State()
    call_time = State()
    comment = State()
    confirm = State()


class SourceForm(StatesGroup):
    name = State()


class Broadcast(StatesGroup):
    waiting = State()
    confirm = State()


# Порядок шагов анкеты и соответствие шаг -> состояние.
STEPS = [
    "name", "phone", "tg", "max", "wa",
    "gender", "medical", "age", "city", "unit", "served", "call_time", "comment",
]
# Шаги, которые задаются не всем: шаг -> условие по уже введённым данным.
CONDITIONAL_STEPS = {"medical": lambda data: data.get("gender") == GENDERS[1]}


def active_steps(data: dict) -> list[str]:
    """Шаги анкеты для конкретного человека (без неподходящих по условию)."""
    return [s for s in STEPS if s not in CONDITIONAL_STEPS or CONDITIONAL_STEPS[s](data)]
STEP_STATES = {step: getattr(Form, step) for step in STEPS}
CONTACT_KEYS = ["phone", "tg", "max", "wa"]
LAST_CONTACT_STEP = "wa"

FIELDS = [
    ("name", "👤", "Имя"),
    ("phone", "📞", "Телефон"),
    ("tg", "✈️", "Telegram"),
    ("max", "💬", "MAX"),
    ("wa", "🟢", "WhatsApp"),
    ("gender", "⚥", "Пол"),
    ("medical", "🩺", "Мед. образование"),
    ("age", "🎂", "Возраст"),
    ("city", "📍", "Город"),
    ("unit", "🎖", "Подразделение"),
    ("served", "🪖", "Служба ранее"),
    ("call_time", "🕒", "Удобное время для звонка"),
    ("comment", "💭", "Комментарий"),
]

GREETING = (
    "👋 <b>Здравствуйте, оставьте заявку и мы свяжемся с вами!</b>\n\n"
    "Это займёт пару минут — несколько простых вопросов.\n\n"
    "<i>Нажимая «Оставить заявку», вы соглашаетесь на обработку персональных данных.</i>"
)

PROMPTS = {
    "name": "👤 <b>Как к вам обращаться?</b>\n\nНапишите ваше имя.",
    "phone": (
        "📞 <b>Номер телефона</b>\n\n"
        "Отправьте номер кнопкой ниже или напишите его вручную.\n"
        "<i>Достаточно указать любой один способ связи — остальные можно пропустить.</i>"
    ),
    "tg": "✈️ <b>Telegram</b>\n\nНапишите ваш @ник или ссылку на профиль.",
    "max": "💬 <b>MAX</b>\n\nУкажите номер телефона или ссылку на ваш профиль в MAX.",
    "wa": "🟢 <b>WhatsApp</b>\n\nУкажите номер, на котором есть WhatsApp.",
    "gender": "⚥ <b>Ваш пол</b>\n\nВыберите вариант.",
    "medical": (
        "🩺 <b>Есть ли у вас медицинское образование?</b>\n\n"
        "Выберите вариант или напишите подробнее — например, «фельдшер» или «медсестра»."
    ),
    "age": "🎂 <b>Сколько вам полных лет?</b>\n\nНапишите число, например: <code>27</code>",
    "city": "📍 <b>Город или регион, где вы находитесь</b>\n\nНапример: <code>Самара</code>",
    "unit": (
        "🎖 <b>Желаемое подразделение</b>\n\n"
        "Выберите из списка или напишите своё. "
        "Если ещё не определились — выберите последний пункт, менеджер подскажет."
    ),
    "served": "🪖 <b>Служили ли вы ранее?</b>\n\nВыберите вариант или напишите подробнее.",
    "call_time": "🕒 <b>Когда вам удобно принять звонок?</b>\n\nВыберите вариант или напишите своё время.",
    "comment": (
        "💭 <b>Хотите что-то добавить?</b>\n\n"
        "Вопросы, пожелания, особенности — всё, что поможет менеджеру. Можно пропустить."
    ),
}

ERRORS = {
    "name": "Пожалуйста, напишите имя текстом (до 60 символов).",
    "phone": "Не похоже на номер телефона. Пример: <code>+7 999 123-45-67</code>",
    "tg": "Не получилось разобрать ник. Пример: <code>@username</code>",
    "max": "Укажите номер телефона или ссылку/ник в MAX.",
    "wa": "Не похоже на номер телефона. Пример: <code>+7 999 123-45-67</code>",
    "gender": "Выберите вариант кнопкой под вопросом: «Мужской» или «Женский».",
    "medical": "Слишком длинно — сократите, пожалуйста, до 100 символов.",
    "age": (
        f"Служба по контракту доступна от {MIN_AGE} до {MAX_AGE} лет. "
        "Проверьте, пожалуйста, возраст и напишите число."
    ),
    "city": "Напишите название города или региона (от 2 до 100 символов).",
    "unit": "Выберите подразделение из списка или напишите его название (до 100 символов).",
    "served": "Слишком длинно — сократите, пожалуйста, до 200 символов.",
    "call_time": "Слишком длинно — сократите, пожалуйста, до 100 символов.",
    "comment": "Слишком длинно — сократите, пожалуйста, до 500 символов.",
}

PENDING_REVIEW = (
    "⏳ <b>Ваша заявка на рассмотрении</b>\n\n"
    "Мы получили вашу заявку от {when} и скоро свяжемся с вами. "
    "Пожалуйста, подождите — подавать заявку повторно не нужно."
)
PENDING_WORK = (
    "🔧 <b>Ваша заявка уже в работе</b>\n\n"
    "Менеджер занимается вашей заявкой от {when} и свяжется с вами. "
    "Пожалуйста, подождите — подавать заявку повторно не нужно."
)

BTN_SHARE = "📱 Поделиться номером"
BTN_SKIP = "⏭ Пропустить"
BTN_BACK = "⬅️ Назад"

bot_props = DefaultBotProperties(parse_mode=ParseMode.HTML)
dp = Dispatcher()

# Два роутера: анкета и админ-команды работают только в личке с ботом,
# а в группе заявок бот реагирует лишь на кнопки статусов и заметки менеджеров.
leads = Router(name="leads")
private = Router(name="private")
private.message.filter(F.chat.type == ChatType.PRIVATE)
private.callback_query.filter(F.message.chat.type == ChatType.PRIVATE)
dp.include_router(leads)
dp.include_router(private)


# ---------- время ----------

def local_dt(ts: str | None) -> datetime | None:
    """Время из SQLite (UTC) -> локальное время."""
    if not ts:
        return None
    try:
        return datetime.strptime(ts, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc).astimezone(LOCAL_TZ)
    except ValueError:
        return None


def fmt_ts(ts: str | None, pattern: str = "%d.%m %H:%M") -> str:
    dt = local_dt(ts)
    return dt.strftime(pattern) if dt else ""


# ---------- база ----------

def ensure_columns(db: sqlite3.Connection, table: str, columns: dict[str, str]) -> None:
    existing = {row[1] for row in db.execute(f"PRAGMA table_info({table})")}
    for name, ddl in columns.items():
        if name not in existing:
            db.execute(f"ALTER TABLE {table} ADD COLUMN {name} {ddl}")


def init_db() -> None:
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS applications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                user_id INTEGER,
                username TEXT,
                full_name TEXT,
                name TEXT,
                phone TEXT,
                telegram TEXT,
                max_contact TEXT,
                whatsapp TEXT
            )
            """
        )
        # Воронка для статистики: одна строка на пользователя, у каждого этапа
        # (запустил бота / начал заявку / завершил заявку) — своя отметка времени,
        # выставляется один раз при первом достижении этапа.
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS funnel (
                user_id INTEGER PRIMARY KEY,
                username TEXT,
                full_name TEXT,
                started_at TEXT,
                started_form_at TEXT,
                completed_at TEXT
            )
            """
        )
        # Миграции: столбцы, появившиеся после первых версий. Старые заявки
        # получают статус «Новая», остальные новые поля у них пустые.
        had_legacy_flag = "legacy" in {
            row[1] for row in db.execute("PRAGMA table_info(applications)")
        }
        ensure_columns(db, "applications", {
            "legacy": "INTEGER DEFAULT 0",
            "gender": "TEXT",
            "medical": "TEXT",
            "age": "INTEGER",
            "city": "TEXT",
            "unit": "TEXT",
            "served": "TEXT",
            "call_time": "TEXT",
            "comment": "TEXT",
            "source": "TEXT",
            "status": "TEXT DEFAULT 'new'",
            "status_by": "TEXT",
            "status_by_id": "INTEGER",
            "status_at": "TEXT",
            "reminded_at": "TEXT",
            "remind_count": "INTEGER DEFAULT 0",
        })
        ensure_columns(db, "funnel", {"blocked_at": "TEXT", "source": "TEXT"})
        # Карточки заявок в группе (по ним ищем заявку при нажатии кнопок и ответах).
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS lead_messages (
                app_id INTEGER NOT NULL,
                chat_id INTEGER NOT NULL,
                message_id INTEGER NOT NULL,
                PRIMARY KEY (chat_id, message_id)
            )
            """
        )
        db.execute("CREATE INDEX IF NOT EXISTS idx_lead_messages_app ON lead_messages (app_id)")
        if not had_legacy_flag:
            # Заявки, поданные до появления карточек со статусами, обработать кнопками
            # нельзя — они не должны блокировать новые заявки этих людей навсегда.
            db.execute(
                "UPDATE applications SET legacy = 1"
                " WHERE id NOT IN (SELECT app_id FROM lead_messages)"
            )
        # Рекламные ссылки: случайный код в ссылке -> понятное название для админа.
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS sources (
                code TEXT PRIMARY KEY,
                name TEXT NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                created_by INTEGER
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS lead_notes (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                app_id INTEGER NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                by_id INTEGER,
                by_name TEXT,
                text TEXT
            )
            """
        )
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS status_history (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                app_id INTEGER NOT NULL,
                created_at TEXT DEFAULT CURRENT_TIMESTAMP,
                status TEXT,
                by_id INTEGER,
                by_name TEXT
            )
            """
        )


def touch_funnel(user: User, stage: str, source: str | None = None) -> None:
    assert stage in ("started_at", "started_form_at", "completed_at")
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        # Источник (метка рекламы) запоминается по первому заходу с меткой.
        db.execute(
            "INSERT INTO funnel (user_id, username, full_name, source) VALUES (?, ?, ?, ?)"
            " ON CONFLICT(user_id) DO UPDATE SET username = excluded.username,"
            " full_name = excluded.full_name,"
            " source = COALESCE(funnel.source, excluded.source)",
            (user.id, user.username, user.full_name, source),
        )
        db.execute(
            f"UPDATE funnel SET {stage} = COALESCE({stage}, CURRENT_TIMESTAMP) WHERE user_id = ?",
            (user.id,),
        )


def had_started_before(user_id: int) -> bool:
    with closing(sqlite3.connect(DB_PATH)) as db:
        row = db.execute(
            "SELECT started_at FROM funnel WHERE user_id = ?", (user_id,)
        ).fetchone()
    return bool(row and row[0])


def get_stats() -> dict:
    with closing(sqlite3.connect(DB_PATH)) as db:
        row = db.execute(
            "SELECT"
            " COUNT(*) FILTER (WHERE started_at IS NOT NULL),"
            " COUNT(*) FILTER (WHERE started_form_at IS NOT NULL),"
            " COUNT(*) FILTER (WHERE completed_at IS NOT NULL),"
            " COUNT(*) FILTER (WHERE blocked_at IS NOT NULL)"
            " FROM funnel"
        ).fetchone()
        total_apps = db.execute("SELECT COUNT(*) FROM applications").fetchone()[0]
        by_status = dict(db.execute(
            "SELECT COALESCE(status, 'new'), COUNT(*) FROM applications GROUP BY 1"
        ).fetchall())
        sources = db.execute(
            "SELECT COALESCE(source, ''),"
            " COUNT(*) FILTER (WHERE started_at IS NOT NULL),"
            " COUNT(*) FILTER (WHERE started_form_at IS NOT NULL),"
            " COUNT(*) FILTER (WHERE completed_at IS NOT NULL)"
            " FROM funnel GROUP BY 1 ORDER BY 2 DESC LIMIT 10"
        ).fetchall()
    started, started_form, completed, blocked = row
    return {
        "started": started,
        "started_form": started_form,
        "completed": completed,
        "blocked": blocked,
        "total_apps": total_apps,
        "by_status": by_status,
        "sources": sources,
    }


def get_broadcast_recipients() -> list[int]:
    """Все, кто хоть раз нажимал /start и не заблокировал бота."""
    with closing(sqlite3.connect(DB_PATH)) as db:
        rows = db.execute(
            "SELECT user_id FROM funnel WHERE started_at IS NOT NULL AND blocked_at IS NULL"
        ).fetchall()
    return [r[0] for r in rows]


def mark_blocked(user_id: int) -> None:
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        db.execute(
            "UPDATE funnel SET blocked_at = CURRENT_TIMESTAMP WHERE user_id = ?", (user_id,)
        )


def source_names() -> dict[str, str]:
    with closing(sqlite3.connect(DB_PATH)) as db:
        return dict(db.execute("SELECT code, name FROM sources").fetchall())


def source_label(code: str | None) -> str | None:
    """Название источника по коду из ссылки (неизвестный код показываем как есть)."""
    return source_names().get(code, code) if code else None


def source_name_taken(name: str) -> bool:
    return any(n.casefold() == name.casefold() for n in source_names().values())


def create_source(name: str, user_id: int) -> str:
    alphabet = string.ascii_lowercase + string.digits
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        while True:
            code = "".join(secrets.choice(alphabet) for _ in range(SOURCE_CODE_LEN))
            if not db.execute("SELECT 1 FROM sources WHERE code = ?", (code,)).fetchone():
                break
        db.execute(
            "INSERT INTO sources (code, name, created_by) VALUES (?, ?, ?)", (code, name, user_id)
        )
    return code


def list_sources(limit: int = 20) -> tuple[int, list[tuple]]:
    """Последние созданные ссылки с числами: запустили / начали заявку / подали."""
    with closing(sqlite3.connect(DB_PATH)) as db:
        total = db.execute("SELECT COUNT(*) FROM sources").fetchone()[0]
        rows = db.execute(
            "SELECT s.code, s.name,"
            " COUNT(f.user_id) FILTER (WHERE f.started_at IS NOT NULL),"
            " COUNT(f.user_id) FILTER (WHERE f.started_form_at IS NOT NULL),"
            " COUNT(f.user_id) FILTER (WHERE f.completed_at IS NOT NULL)"
            " FROM sources s LEFT JOIN funnel f ON f.source = s.code"
            " GROUP BY s.code ORDER BY s.rowid DESC LIMIT ?",
            (limit,),
        ).fetchall()
    return total, rows


def save_application(user: User, data: dict) -> int:
    age = data.get("age")
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        row = db.execute("SELECT source FROM funnel WHERE user_id = ?", (user.id,)).fetchone()
        cur = db.execute(
            "INSERT INTO applications"
            " (user_id, username, full_name, name, phone, telegram, max_contact, whatsapp,"
            "  gender, medical, age, city, unit, served, call_time, comment, source)"
            " VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (
                user.id, user.username, user.full_name,
                data.get("name"), data.get("phone"), data.get("tg"),
                data.get("max"), data.get("wa"),
                data.get("gender"), data.get("medical"), int(age) if age else None, data.get("city"),
                data.get("unit"), data.get("served"), data.get("call_time"),
                data.get("comment"), row[0] if row else None,
            ),
        )
        return cur.lastrowid


def active_application(user_id: int) -> dict | None:
    """Заявка человека, которая ещё ждёт обработки (блокирует подачу новой)."""
    marks = ",".join("?" * len(BLOCKING_STATUSES))
    with closing(sqlite3.connect(DB_PATH)) as db:
        db.row_factory = sqlite3.Row
        row = db.execute(
            "SELECT id, created_at, COALESCE(status, 'new') AS status FROM applications"
            " WHERE user_id = ? AND COALESCE(legacy, 0) = 0"
            f" AND COALESCE(status, 'new') IN ({marks}) ORDER BY id DESC LIMIT 1",
            (user_id, *BLOCKING_STATUSES),
        ).fetchone()
    return dict(row) if row else None


def get_application(app_id: int) -> dict | None:
    with closing(sqlite3.connect(DB_PATH)) as db:
        db.row_factory = sqlite3.Row
        row = db.execute("SELECT * FROM applications WHERE id = ?", (app_id,)).fetchone()
        if row is None:
            return None
        notes = db.execute(
            "SELECT by_name, text, created_at FROM lead_notes WHERE app_id = ? ORDER BY id",
            (app_id,),
        ).fetchall()
    app = dict(row)
    app["status"] = app["status"] or "new"
    app["notes"] = [dict(n) for n in notes]
    return app


def save_lead_message(app_id: int, chat_id: int, message_id: int) -> None:
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        db.execute(
            "INSERT OR REPLACE INTO lead_messages (app_id, chat_id, message_id) VALUES (?, ?, ?)",
            (app_id, chat_id, message_id),
        )


def lead_message_refs(app_id: int) -> list[tuple[int, int]]:
    with closing(sqlite3.connect(DB_PATH)) as db:
        return db.execute(
            "SELECT chat_id, message_id FROM lead_messages WHERE app_id = ?", (app_id,)
        ).fetchall()


def find_app_by_message(chat_id: int, message_id: int) -> int | None:
    with closing(sqlite3.connect(DB_PATH)) as db:
        row = db.execute(
            "SELECT app_id FROM lead_messages WHERE chat_id = ? AND message_id = ?",
            (chat_id, message_id),
        ).fetchone()
    return row[0] if row else None


def set_status(app_id: int, status: str, user: User) -> bool | None:
    """True — статус изменён, False — уже был таким, None — заявки нет."""
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        row = db.execute(
            "SELECT COALESCE(status, 'new') FROM applications WHERE id = ?", (app_id,)
        ).fetchone()
        if row is None:
            return None
        if row[0] == status:
            return False
        db.execute(
            "UPDATE applications SET status = ?, status_by = ?, status_by_id = ?,"
            " status_at = CURRENT_TIMESTAMP, remind_count = 0, reminded_at = NULL WHERE id = ?",
            (status, user.full_name, user.id, app_id),
        )
        db.execute(
            "INSERT INTO status_history (app_id, status, by_id, by_name) VALUES (?, ?, ?, ?)",
            (app_id, status, user.id, user.full_name),
        )
    return True


def add_note(app_id: int, user: User, text: str) -> None:
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        db.execute(
            "INSERT INTO lead_notes (app_id, by_id, by_name, text) VALUES (?, ?, ?, ?)",
            (app_id, user.id, user.full_name, text),
        )


def due_reminders() -> list[tuple[int, int, int, int]]:
    """Заявки со статусом «Новая», которые ждут дольше REMIND_AFTER_MIN минут."""
    with closing(sqlite3.connect(DB_PATH)) as db:
        return db.execute(
            "SELECT a.id, m.chat_id, m.message_id,"
            " CAST((julianday('now') - julianday(a.created_at)) * 1440 AS INTEGER)"
            " FROM applications a JOIN lead_messages m ON m.app_id = a.id"
            " WHERE COALESCE(a.status, 'new') = 'new'"
            " AND COALESCE(a.remind_count, 0) < ?"
            " AND COALESCE(a.reminded_at, a.created_at) <= datetime('now', ?)",
            (REMIND_MAX, f"-{REMIND_AFTER_MIN} minutes"),
        ).fetchall()


def mark_reminded(app_ids: set[int]) -> None:
    with closing(sqlite3.connect(DB_PATH)) as db, db:
        db.executemany(
            "UPDATE applications SET remind_count = COALESCE(remind_count, 0) + 1,"
            " reminded_at = CURRENT_TIMESTAMP WHERE id = ?",
            [(i,) for i in app_ids],
        )


def safe_cell(value):
    """Строки, начинающиеся с «=», Excel считал бы формулой — экранируем."""
    if isinstance(value, str) and value.startswith("="):
        return "'" + value
    return value


EXPORT_COLUMNS = [
    ("#", 5), ("Дата и время", 17), ("Имя", 16), ("Телефон", 16), ("Telegram", 14),
    ("MAX", 14), ("WhatsApp", 14), ("Пол", 10), ("Мед. образование", 16), ("Возраст", 9), ("Город", 16),
    ("Подразделение", 20), ("Служба ранее", 18), ("Удобное время", 16),
    ("Комментарий клиента", 30), ("Источник", 14), ("Статус", 16), ("Менеджер", 16),
    ("Статус изменён", 17), ("Заметки менеджера", 40), ("Профиль", 20),
    ("Username", 14), ("User ID", 12),
]


def build_export_xlsx() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Заявки"
    ws.append([name for name, _ in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    with closing(sqlite3.connect(DB_PATH)) as db:
        db.row_factory = sqlite3.Row
        apps = db.execute("SELECT * FROM applications ORDER BY id").fetchall()
        names = source_names()
        notes: dict[int, list[str]] = {}
        for n in db.execute(
            "SELECT app_id, by_name, text, created_at FROM lead_notes ORDER BY id"
        ):
            notes.setdefault(n["app_id"], []).append(
                f"{n['by_name']} ({fmt_ts(n['created_at'])}): {n['text']}"
            )
    for a in apps:
        status = a["status"] or "new"
        ws.append([safe_cell(v) for v in [
            a["id"], fmt_ts(a["created_at"], "%d.%m.%Y %H:%M"), a["name"], a["phone"],
            a["telegram"], a["max_contact"], a["whatsapp"], a["gender"], a["medical"], a["age"], a["city"],
            a["unit"], a["served"], a["call_time"], a["comment"], names.get(a["source"], a["source"]),
            STATUSES.get(status, STATUSES["new"])[1],
            a["status_by"] if status != "new" else "",
            fmt_ts(a["status_at"], "%d.%m.%Y %H:%M") if status != "new" else "",
            "\n".join(notes.get(a["id"], [])),
            a["full_name"], f"@{a['username']}" if a["username"] else "", a["user_id"],
        ]])
    for i, (_, width) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------- проверка введённых значений ----------

def norm_phone(text: str) -> str | None:
    digits = re.sub(r"\D", "", text)
    if len(digits) == 10:
        digits = "7" + digits
    elif len(digits) == 11 and digits[0] == "8":
        digits = "7" + digits[1:]
    if not 10 <= len(digits) <= 15:
        return None
    return "+" + digits


def norm_tg(text: str) -> str | None:
    m = re.fullmatch(
        r"(?:https?://)?(?:t\.me/|telegram\.me/)?@?([A-Za-z][A-Za-z0-9_]{4,31})/?", text.strip()
    )
    if m:
        return "@" + m.group(1)
    return norm_phone(text)


def norm_max(text: str) -> str | None:
    phone = norm_phone(text)
    if phone:
        return phone
    return text.strip() if 3 <= len(text.strip()) <= 100 else None


def norm_wa(text: str) -> str | None:
    m = re.search(r"wa\.me/\+?(\d{10,15})", text)
    if m:
        return "+" + m.group(1)
    return norm_phone(text)


def norm_name(text: str) -> str | None:
    text = text.strip()
    return text if 1 <= len(text) <= 60 else None


def norm_gender(text: str) -> str | None:
    t = text.strip().lower()
    if t in ("м", "муж", "мужской", "мужчина"):
        return GENDERS[0]
    if t in ("ж", "жен", "женский", "женщина"):
        return GENDERS[1]
    return None


def norm_age(text: str) -> str | None:
    m = re.fullmatch(r"\D*(\d{1,3})\D*", text.strip())
    if not m:
        return None
    age = int(m.group(1))
    return str(age) if MIN_AGE <= age <= MAX_AGE else None


def text_normalizer(min_len: int, max_len: int):
    def norm(text: str) -> str | None:
        text = text.strip()
        return text if min_len <= len(text) <= max_len else None
    return norm


NORMALIZERS = {
    "name": norm_name,
    "phone": norm_phone,
    "tg": norm_tg,
    "max": norm_max,
    "wa": norm_wa,
    "gender": norm_gender,
    "medical": text_normalizer(1, 100),
    "age": norm_age,
    "city": text_normalizer(2, 100),
    "unit": text_normalizer(2, 100),
    "served": text_normalizer(1, 200),
    "call_time": text_normalizer(1, 100),
    "comment": text_normalizer(1, 500),
}


def parse_source(arg: str | None) -> str | None:
    """Метка рекламы из ссылки t.me/бот?start=метка."""
    return arg if arg and re.fullmatch(r"[A-Za-z0-9_-]{1,64}", arg) else None


# ---------- оформление ----------

def header(step: str, data: dict) -> str:
    steps = active_steps(data)
    i = steps.index(step)
    bar = "●" * (i + 1) + "○" * (len(steps) - i - 1)
    return f"<b>Шаг {i + 1} из {len(steps)}</b>  {bar}"


def card(data: dict, skip_empty: bool = False) -> str:
    lines = []
    for key, icon, label in FIELDS:
        if key in CONDITIONAL_STEPS and not CONDITIONAL_STEPS[key](data):
            continue
        value = data.get(key)
        if value:
            lines.append(f"{icon} <b>{label}:</b> {escape(str(value))}")
        elif not skip_empty:
            lines.append(f"{icon} <b>{label}:</b> —")
    return "\n".join(lines)


def btn(text: str, cb: str) -> InlineKeyboardButton:
    return InlineKeyboardButton(text=text, callback_data=cb)


def choice_rows(step: str) -> list[list[InlineKeyboardButton]]:
    """Варианты ответа: короткие по два в ряд, длинные — каждый в своей строке."""
    rows, pair = [], []
    for n, label in enumerate(CHOICES[step]):
        b = btn(label, f"pick:{n}")
        if len(label) > 22:
            if pair:
                rows.append(pair)
                pair = []
            rows.append([b])
        else:
            pair.append(b)
            if len(pair) == 2:
                rows.append(pair)
                pair = []
    if pair:
        rows.append(pair)
    return rows


def step_keyboard(step: str, data: dict, user: User) -> InlineKeyboardMarkup:
    i = active_steps(data).index(step)
    rows = []
    if step == "name" and user.first_name:
        rows.append([btn(f"👋 Меня зовут {user.first_name}", "use:name")])
    if step == "tg" and user.username:
        rows.append([btn(f"✈️ Мой @{user.username}", "use:tg")])
    if step in ("max", "wa") and data.get("phone"):
        rows.append([btn(f"📞 Тот же номер {data['phone']}", "use:phone")])
    if step in CHOICES:
        rows.extend(choice_rows(step))
    nav = []
    if i > 0:
        nav.append(btn(BTN_BACK, "back"))
    if step not in REQUIRED_STEPS:
        nav.append(btn(BTN_SKIP, "skip"))
    if nav:
        rows.append(nav)
    rows.append([btn("✖️ Отмена", "cancel")])
    return InlineKeyboardMarkup(inline_keyboard=rows)


def phone_keyboard() -> ReplyKeyboardMarkup:
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text=BTN_SHARE, request_contact=True)],
            [KeyboardButton(text=BTN_BACK), KeyboardButton(text=BTN_SKIP)],
        ],
        resize_keyboard=True,
        input_field_placeholder="+7 999 123-45-67",
    )


def confirm_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [btn("✅ Отправить заявку", "send")],
            [btn("✏️ Заполнить заново", "restart"), btn(BTN_BACK, "back")],
            [btn("✖️ Отмена", "cancel")],
        ]
    )


def apply_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[btn("📝 Оставить заявку", "apply")]])


def broadcast_confirm_keyboard() -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(
        inline_keyboard=[
            [btn("📣 Разослать", "bc_send")],
            [btn("✖️ Отмена", "bc_cancel")],
        ]
    )


# ---------- карточка заявки в группе ----------

def pending_notice(app: dict) -> str:
    template = PENDING_WORK if app["status"] == "work" else PENDING_REVIEW
    return template.format(when=fmt_ts(app["created_at"]))


def lead_text(app: dict) -> str:
    status = app["status"]
    icon, label = STATUSES.get(status, STATUSES["new"])
    data = {
        "name": app["name"], "phone": app["phone"], "tg": app["telegram"],
        "max": app["max_contact"], "wa": app["whatsapp"], "gender": app["gender"],
        "medical": app["medical"], "age": app["age"], "city": app["city"], "unit": app["unit"],
        "served": app["served"], "call_time": app["call_time"], "comment": app["comment"],
    }
    username = f"@{app['username']}" if app["username"] else "нет username"
    lines = [
        f"📋 <b>Заявка #{app['id']}</b> · {icon} {label}",
        "",
        card(data, skip_empty=True),
        "",
        f"🔗 <a href=\"tg://user?id={app['user_id']}\">{escape(app['full_name'] or '—')}</a>"
        f" · {escape(username)} · ID <code>{app['user_id']}</code>",
    ]
    if app["source"]:
        lines.append(f"🔖 Источник: <b>{escape(source_label(app['source']))}</b>")
    if status != "new":
        lines += [
            "",
            f"{icon} <b>{label}</b> — {escape(app['status_by'] or '—')} · {fmt_ts(app['status_at'])}",
        ]
    notes = app["notes"][-5:]
    if notes:
        lines += ["", "📝 <b>Заметки:</b>"]
        for n in notes:
            lines.append(
                f"• <b>{escape(n['by_name'] or '—')}</b> ({fmt_ts(n['created_at'])}): {escape(n['text'])}"
            )
    return "\n".join(lines)


def lead_keyboard(app_id: int, current: str) -> InlineKeyboardMarkup:
    buttons = [
        btn(STATUS_BUTTONS[key], f"st:{app_id}:{key}")
        for key in STATUS_BUTTON_ORDER
        if key != current
    ]
    return InlineKeyboardMarkup(
        inline_keyboard=[buttons[i:i + 2] for i in range(0, len(buttons), 2)]
    )


async def send_lead_cards(bot: Bot, app_id: int) -> None:
    app = get_application(app_id)
    text, markup = lead_text(app), lead_keyboard(app_id, app["status"])
    for chat_id in ADMIN_IDS:
        try:
            msg = await bot.send_message(chat_id, text, reply_markup=markup)
            save_lead_message(app_id, chat_id, msg.message_id)
        except Exception:
            logging.exception("Не удалось отправить заявку в чат %s", chat_id)


async def refresh_cards(bot: Bot, app_id: int) -> None:
    app = get_application(app_id)
    if app is None:
        return
    text, markup = lead_text(app), lead_keyboard(app_id, app["status"])
    for chat_id, message_id in lead_message_refs(app_id):
        with suppress(TelegramBadRequest):
            await bot.edit_message_text(
                text, chat_id=chat_id, message_id=message_id, reply_markup=markup
            )


def can_manage(user_id: int) -> bool:
    return not MANAGER_IDS or user_id in MANAGER_IDS or user_id in ADMIN_USER_IDS


# ---------- движок анкеты ----------

async def clear_prev(bot: Bot, chat_id: int, data: dict) -> None:
    """Убирает кнопки с предыдущего вопроса (и клавиатуру с номером телефона)."""
    if data.get("prompt_id"):
        with suppress(TelegramBadRequest):
            await bot.edit_message_reply_markup(
                chat_id=chat_id, message_id=data["prompt_id"], reply_markup=None
            )
    if data.get("reply_kb"):
        with suppress(TelegramBadRequest):
            tmp = await bot.send_message(chat_id, "…", reply_markup=ReplyKeyboardRemove())
            await tmp.delete()


async def show_step(bot: Bot, chat_id: int, state: FSMContext, step: str, user: User) -> None:
    data = await state.get_data()
    await clear_prev(bot, chat_id, data)
    text = f"{header(step, data)}\n\n{PROMPTS[step]}"
    markup = phone_keyboard() if step == "phone" else step_keyboard(step, data, user)
    msg = await bot.send_message(chat_id, text, reply_markup=markup)
    await state.set_state(STEP_STATES[step])
    await state.update_data(step=step, prompt_id=msg.message_id, reply_kb=step == "phone")


async def show_confirm(bot: Bot, chat_id: int, state: FSMContext) -> None:
    data = await state.get_data()
    await clear_prev(bot, chat_id, data)
    text = f"📋 <b>Проверьте вашу заявку</b>\n\n{card(data)}\n\nВсё верно?"
    msg = await bot.send_message(chat_id, text, reply_markup=confirm_keyboard())
    await state.set_state(Form.confirm)
    await state.update_data(step="confirm", prompt_id=msg.message_id, reply_kb=False)


async def save_and_next(
    bot: Bot, chat_id: int, state: FSMContext, user: User, step: str, value: str | None
) -> None:
    await state.update_data(**{step: value})
    if step == "gender" and value != GENDERS[1]:
        await state.update_data(medical=None)  # вопрос только для женщин
    data = await state.get_data()
    if step == LAST_CONTACT_STEP:
        if not any(data.get(k) for k in CONTACT_KEYS):
            await bot.send_message(
                chat_id,
                "⚠️ Нужен хотя бы один способ связи — укажите телефон, Telegram, MAX или WhatsApp.",
            )
            await show_step(bot, chat_id, state, "phone", user)
            return
    steps = active_steps(data)
    i = steps.index(step)
    if i + 1 < len(steps):
        await show_step(bot, chat_id, state, steps[i + 1], user)
    else:
        await show_confirm(bot, chat_id, state)


async def go_back(bot: Bot, chat_id: int, state: FSMContext, user: User, step: str) -> None:
    steps = active_steps(await state.get_data())
    prev = steps[-1] if step == "confirm" else steps[max(steps.index(step) - 1, 0)]
    await show_step(bot, chat_id, state, prev, user)


async def start_form(bot: Bot, chat_id: int, state: FSMContext, user: User) -> None:
    data = await state.get_data()
    await clear_prev(bot, chat_id, data)
    await state.clear()
    await show_step(bot, chat_id, state, "name", user)


async def guard(cb: CallbackQuery, state: FSMContext) -> dict | None:
    """Пропускает только нажатия на кнопки актуального сообщения анкеты."""
    data = await state.get_data()
    if cb.message is None or cb.message.message_id != data.get("prompt_id"):
        await cb.answer("Эта кнопка уже неактуальна")
        return None
    return data


# ---------- группа заявок: статусы и заметки ----------

@leads.callback_query(F.data.startswith("st:"), F.message.chat.id.in_(ADMIN_CHAT_SET))
async def on_status(cb: CallbackQuery, bot: Bot) -> None:
    if not can_manage(cb.from_user.id):
        await cb.answer("У вас нет прав менять статусы", show_alert=True)
        return
    try:
        _, app_id_str, status = cb.data.split(":")
        app_id = int(app_id_str)
    except ValueError:
        await cb.answer()
        return
    if status not in STATUSES:
        await cb.answer()
        return
    changed = set_status(app_id, status, cb.from_user)
    if changed is None:
        await cb.answer("Заявка не найдена", show_alert=True)
        return
    icon, label = STATUSES[status]
    await cb.answer(f"{icon} {label}" if changed else "Этот статус уже стоит")
    await refresh_cards(bot, app_id)


@leads.message(
    F.chat.id.in_(ADMIN_CHAT_SET), F.reply_to_message, F.text, ~F.text.startswith("/")
)
async def on_note(message: Message, bot: Bot) -> None:
    """Ответ менеджера на карточку заявки сохраняется как заметка к ней."""
    if message.from_user is None or not can_manage(message.from_user.id):
        return
    app_id = find_app_by_message(message.chat.id, message.reply_to_message.message_id)
    if app_id is None:
        return
    add_note(app_id, message.from_user, message.text.strip()[:300])
    await refresh_cards(bot, app_id)
    with suppress(Exception):
        await message.react([ReactionTypeEmoji(emoji="👍")])


def fmt_wait(minutes: int) -> str:
    return f"{minutes // 60} ч {minutes % 60:02d} мин" if minutes >= 60 else f"{minutes} мин"


async def send_reminders(bot: Bot) -> None:
    done: set[int] = set()
    for app_id, chat_id, message_id, waited in due_reminders():
        try:
            await bot.send_message(
                chat_id,
                f"⏰ <b>Заявка #{app_id}</b> без обработки уже {fmt_wait(waited)}. "
                "Нажмите «Беру в работу» под карточкой.",
                reply_parameters=ReplyParameters(
                    message_id=message_id, allow_sending_without_reply=True
                ),
            )
        except Exception:
            logging.exception("Не удалось отправить напоминание по заявке %s", app_id)
        done.add(app_id)
    if done:
        mark_reminded(done)


async def reminder_loop(bot: Bot) -> None:
    while True:
        try:
            await send_reminders(bot)
        except Exception:
            logging.exception("Ошибка проверки напоминаний")
        await asyncio.sleep(60)


# ---------- личка: анкета и админ-команды ----------

async def notify_start(
    bot: Bot, user: User, is_new: bool, source: str | None, pending: dict | None = None
) -> None:
    label = "🆕 Новый пользователь запустил бота" if is_new else "🔁 Пользователь снова нажал /start"
    username = f"@{user.username}" if user.username else "нет username"
    text = (
        f"▶️ <b>{label}</b>\n\n"
        f"👤 <a href=\"tg://user?id={user.id}\">{escape(user.full_name)}</a> · {escape(username)}\n"
        f"ID: <code>{user.id}</code>"
    )
    if source:
        text += f"\n🔖 Источник: <b>{escape(source_label(source))}</b>"
    if pending:
        text += f"\n⏳ Заявка #{pending['id']} уже ждёт обработки"
    for admin_id in ADMIN_USER_IDS:
        try:
            await bot.send_message(admin_id, text)
        except Exception:
            logging.exception("Не удалось отправить уведомление о /start админу %s", admin_id)


@private.message(CommandStart())
async def on_start(
    message: Message, state: FSMContext, bot: Bot, command: CommandObject
) -> None:
    source = parse_source(command.args)
    is_new = not had_started_before(message.from_user.id)
    touch_funnel(message.from_user, "started_at", source)
    pending = active_application(message.from_user.id)
    await notify_start(bot, message.from_user, is_new, source, pending)
    await clear_prev(bot, message.chat.id, await state.get_data())
    await state.clear()
    if pending:
        await message.answer(pending_notice(pending))
    else:
        await message.answer(GREETING, reply_markup=apply_keyboard())


@private.message(Command("cancel"))
async def on_cancel_cmd(message: Message, state: FSMContext, bot: Bot) -> None:
    current = await state.get_state()
    if current in (Broadcast.waiting.state, Broadcast.confirm.state):
        await state.clear()
        await message.answer("🚫 Рассылка отменена.")
        return
    if current == SourceForm.name.state:
        await state.clear()
        await message.answer("🚫 Создание ссылки отменено.")
        return
    await clear_prev(bot, message.chat.id, await state.get_data())
    await state.clear()
    await message.answer("🚫 Заявка отменена.", reply_markup=apply_keyboard())


def pct(part: int, whole: int) -> str:
    return f"{part / whole:.0%}" if whole else "—"


@private.message(Command("stats"))
async def on_stats(message: Message) -> None:
    if message.from_user.id not in ADMIN_USER_IDS:
        return
    s = get_stats()
    started, started_form, completed, blocked = (
        s["started"], s["started_form"], s["completed"], s["blocked"]
    )
    by_status = s["by_status"]
    workable = s["total_apps"] - by_status.get("junk", 0)
    agreed = by_status.get("agreed", 0)

    lines = [
        "📊 <b>Статистика бота</b>",
        "",
        f"🚀 Запустили бота: <b>{started}</b>",
        f"📝 Начали заполнять заявку: <b>{started_form}</b> ({pct(started_form, started)} от запустивших)",
        f"✅ Завершили заявку: <b>{completed}</b> ({pct(completed, started_form)} от начавших, "
        f"{pct(completed, started)} от запустивших)",
        f"🚫 Заблокировали бота: <b>{blocked}</b> ({pct(blocked, started)} от запустивших)",
        "",
        f"📨 Всего заявок отправлено: <b>{s['total_apps']}</b>",
        "",
        "📋 <b>Заявки по статусам</b>",
    ]
    for key, (icon, label) in STATUSES.items():
        lines.append(f"{icon} {label}: <b>{by_status.get(key, 0)}</b>")
    lines.append(
        f"🎯 Согласились: <b>{agreed}</b> из {workable} ({pct(agreed, workable)}, без учёта мусора)"
    )
    if s["sources"]:
        lines += ["", "🔖 <b>Источники</b> <i>(запустили → начали → подали)</i>"]
        names = source_names()
        for source, st, fm, done in s["sources"]:
            name = escape(names.get(source, source)) if source else "без метки"
            lines.append(f"<b>{name}</b> — {st} → {fm} → {done} ({pct(done, st)})")
    lines += [
        "",
        "<i>Блокировку бот узнаёт только при попытке написать пользователю "
        "(например, во время рассылки), поэтому число может быть занижено.</i>",
    ]
    await message.answer("\n".join(lines))


@private.message(Command("export"))
async def on_export(message: Message) -> None:
    if message.from_user.id not in ADMIN_USER_IDS:
        return
    with closing(sqlite3.connect(DB_PATH)) as db:
        count = db.execute("SELECT COUNT(*) FROM applications").fetchone()[0]
    if not count:
        await message.answer("Заявок пока нет.")
        return
    buf = build_export_xlsx()
    filename = f"leads_{datetime.now(LOCAL_TZ):%Y-%m-%d_%H%M}.xlsx"
    await message.answer_document(
        BufferedInputFile(buf.read(), filename=filename),
        caption=f"📊 Заявки: {count} шт.",
    )


async def links_view(bot: Bot) -> tuple[str, InlineKeyboardMarkup]:
    username = (await bot.me()).username
    total, rows = list_sources()
    lines = [
        "🔗 <b>Ссылки для рекламы</b>",
        "",
        "Каждая ссылка помечает людей, пришедших по ней: в заявках, статистике и Excel "
        "будет видно название источника. Нажмите на ссылку, чтобы скопировать.",
    ]
    if not rows:
        lines += ["", "Пока нет ни одной ссылки — нажмите «Создать ссылку»."]
    for code, name, started, form, done in rows:
        lines += [
            "",
            f"<b>{escape(name)}</b> — {started} → {form} → {done}",
            f"<code>https://t.me/{username}?start={code}</code>",
        ]
    if rows:
        lines += ["", "<i>Числа: запустили → начали заявку → подали.</i>"]
    if total > len(rows):
        lines.append(f"<i>Показаны последние {len(rows)} из {total}.</i>")
    markup = InlineKeyboardMarkup(inline_keyboard=[[btn("➕ Создать ссылку", "lk_new")]])
    return "\n".join(lines), markup


@private.message(Command("links"))
async def on_links(message: Message, state: FSMContext, bot: Bot) -> None:
    if message.from_user.id not in ADMIN_USER_IDS:
        return
    await clear_prev(bot, message.chat.id, await state.get_data())
    await state.clear()
    text, markup = await links_view(bot)
    await message.answer(text, reply_markup=markup)


@private.callback_query(F.data == "lk_list")
async def on_links_list(cb: CallbackQuery, bot: Bot) -> None:
    if cb.from_user.id not in ADMIN_USER_IDS:
        await cb.answer()
        return
    await cb.answer()
    text, markup = await links_view(bot)
    await cb.message.answer(text, reply_markup=markup)


@private.callback_query(F.data == "lk_new")
async def on_link_new(cb: CallbackQuery, state: FSMContext) -> None:
    if cb.from_user.id not in ADMIN_USER_IDS:
        await cb.answer()
        return
    await cb.answer()
    await state.set_state(SourceForm.name)
    await cb.message.answer(
        "✍️ <b>Название источника</b>\n\n"
        "Напишите, как назвать эту ссылку, например: <code>залив 1</code>. "
        "Название видите только вы.\n\n/cancel — отмена"
    )


@private.message(StateFilter(SourceForm.name), F.text, ~F.text.startswith("/"))
async def on_link_name(message: Message, state: FSMContext, bot: Bot) -> None:
    name = " ".join(message.text.split())
    if not 1 <= len(name) <= 50:
        await message.answer("Название должно быть от 1 до 50 символов. Попробуйте ещё раз.")
        return
    if source_name_taken(name):
        await message.answer("Источник с таким названием уже есть — придумайте другое.")
        return
    code = create_source(name, message.from_user.id)
    await state.clear()
    username = (await bot.me()).username
    await message.answer(
        f"✅ Ссылка для источника <b>{escape(name)}</b> создана:\n\n"
        f"<code>https://t.me/{username}?start={code}</code>\n\n"
        "Вставьте её в рекламу — всех, кто придёт по ней, бот пометит этим названием.",
        reply_markup=InlineKeyboardMarkup(inline_keyboard=[
            [btn("➕ Создать ещё", "lk_new")],
            [btn("📋 Все ссылки", "lk_list")],
        ]),
    )


@private.message(StateFilter(SourceForm.name))
async def on_link_name_other(message: Message) -> None:
    await message.answer("Отправьте название текстом, например: <code>залив 1</code>.")


@private.message(Command("broadcast"))
async def on_broadcast_cmd(message: Message, state: FSMContext) -> None:
    if message.from_user.id not in ADMIN_USER_IDS:
        return
    recipients = get_broadcast_recipients()
    if not recipients:
        await message.answer("Пока нет ни одного пользователя, который запускал бота.")
        return
    await state.set_state(Broadcast.waiting)
    await message.answer(
        f"📣 <b>Рассылка</b> ({len(recipients)} получателей)\n\n"
        "Отправьте сообщение, которое разослать — текст, фото, видео или документ. "
        "Оно уйдёт пользователям в точности так, как вы его отправите.\n\n"
        "/cancel — чтобы отменить."
    )


@private.message(StateFilter(Broadcast.waiting))
async def on_broadcast_content(message: Message, state: FSMContext) -> None:
    recipients = get_broadcast_recipients()
    if not recipients:
        await state.clear()
        await message.answer("Пока нет ни одного пользователя, который запускал бота.")
        return
    await state.update_data(from_chat_id=message.chat.id, message_id=message.message_id)
    await state.set_state(Broadcast.confirm)
    await message.answer(
        f"👆 Сообщение выше уйдёт <b>{len(recipients)}</b> пользователям. Разослать?",
        reply_markup=broadcast_confirm_keyboard(),
    )


@private.callback_query(StateFilter(Broadcast.confirm), F.data == "bc_cancel")
async def on_broadcast_cancel(cb: CallbackQuery, state: FSMContext) -> None:
    await state.clear()
    await cb.answer()
    with suppress(TelegramBadRequest):
        await cb.message.edit_reply_markup(reply_markup=None)
    await cb.message.answer("🚫 Рассылка отменена.")


@private.callback_query(StateFilter(Broadcast.confirm), F.data == "bc_send")
async def on_broadcast_send(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    data = await state.get_data()
    await state.clear()
    await cb.answer()
    with suppress(TelegramBadRequest):
        await cb.message.edit_reply_markup(reply_markup=None)

    recipients = get_broadcast_recipients()
    total = len(recipients)
    status = await cb.message.answer(f"⏳ Рассылка: 0/{total}")
    ok = fail = 0

    for i, user_id in enumerate(recipients, start=1):
        try:
            await bot.copy_message(
                chat_id=user_id,
                from_chat_id=data["from_chat_id"],
                message_id=data["message_id"],
            )
            ok += 1
        except TelegramForbiddenError:
            fail += 1
            mark_blocked(user_id)
        except TelegramRetryAfter as e:
            await asyncio.sleep(e.retry_after)
            try:
                await bot.copy_message(
                    chat_id=user_id,
                    from_chat_id=data["from_chat_id"],
                    message_id=data["message_id"],
                )
                ok += 1
            except Exception:
                fail += 1
        except Exception:
            logging.exception("Не удалось разослать сообщение пользователю %s", user_id)
            fail += 1

        if i % 20 == 0 or i == total:
            with suppress(TelegramBadRequest):
                await status.edit_text(f"⏳ Рассылка: {i}/{total} (✅ {ok} · ⚠️ {fail})")
        await asyncio.sleep(0.05)

    await status.edit_text(
        f"✅ <b>Рассылка завершена</b>\n\nДоставлено: {ok}\nНе доставлено: {fail}"
    )


@private.callback_query(F.data == "apply")
async def on_apply(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    pending = active_application(cb.from_user.id)
    if pending:
        await cb.answer("Ваша заявка уже на рассмотрении", show_alert=True)
        with suppress(TelegramBadRequest):
            await cb.message.edit_text(pending_notice(pending), reply_markup=None)
        return
    await cb.answer()
    touch_funnel(cb.from_user, "started_form_at")
    with suppress(TelegramBadRequest):
        await cb.message.edit_reply_markup(reply_markup=None)
    await start_form(bot, cb.message.chat.id, state, cb.from_user)


@private.callback_query(F.data == "cancel")
async def on_cancel(cb: CallbackQuery, state: FSMContext) -> None:
    await cb.answer()
    await state.clear()
    with suppress(TelegramBadRequest):
        await cb.message.edit_text("🚫 Заявка отменена.", reply_markup=apply_keyboard())


@private.callback_query(StateFilter(Form.confirm), F.data == "send")
async def on_send(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    data = await guard(cb, state)
    if data is None:
        return
    await state.clear()
    pending = active_application(cb.from_user.id)
    if pending:
        await cb.answer("Ваша заявка уже на рассмотрении", show_alert=True)
        with suppress(TelegramBadRequest):
            await cb.message.edit_text(pending_notice(pending), reply_markup=None)
        return
    app_id = save_application(cb.from_user, data)
    touch_funnel(cb.from_user, "completed_at")
    await send_lead_cards(bot, app_id)
    name = f", {escape(data['name'])}" if data.get("name") else ""
    await cb.answer("Заявка отправлена!")
    await cb.message.edit_text(
        f"🎉 <b>Заявка отправлена!</b>\n\nСпасибо{name}! Мы свяжемся с вами в ближайшее время.",
        reply_markup=None,
    )


@private.callback_query(StateFilter(Form.confirm), F.data == "restart")
async def on_restart(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    if await guard(cb, state) is None:
        return
    await cb.answer()
    await start_form(bot, cb.message.chat.id, state, cb.from_user)


STEP_FILTER = StateFilter(*STEP_STATES.values(), Form.confirm)


@private.callback_query(STEP_FILTER, F.data == "back")
async def on_back(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    data = await guard(cb, state)
    if data is None:
        return
    await cb.answer()
    await go_back(bot, cb.message.chat.id, state, cb.from_user, data["step"])


@private.callback_query(StateFilter(*STEP_STATES.values()), F.data == "skip")
async def on_skip(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    data = await guard(cb, state)
    if data is None:
        return
    if data["step"] in REQUIRED_STEPS:
        await cb.answer("Этот вопрос обязательный", show_alert=True)
        return
    await cb.answer()
    await save_and_next(bot, cb.message.chat.id, state, cb.from_user, data["step"], None)


@private.callback_query(StateFilter(*STEP_STATES.values()), F.data.startswith("use:"))
async def on_use(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    data = await guard(cb, state)
    if data is None:
        return
    source = cb.data.split(":", 1)[1]
    value = {
        "name": cb.from_user.first_name,
        "tg": f"@{cb.from_user.username}" if cb.from_user.username else None,
        "phone": data.get("phone"),
    }.get(source)
    if not value:
        await cb.answer("Нет данных для подстановки")
        return
    await cb.answer()
    await save_and_next(bot, cb.message.chat.id, state, cb.from_user, data["step"], value)


@private.callback_query(StateFilter(*STEP_STATES.values()), F.data.startswith("pick:"))
async def on_pick(cb: CallbackQuery, state: FSMContext, bot: Bot) -> None:
    data = await guard(cb, state)
    if data is None:
        return
    options = CHOICES.get(data["step"])
    try:
        value = options[int(cb.data.split(":", 1)[1])]
    except (TypeError, ValueError, IndexError):
        await cb.answer("Эта кнопка уже неактуальна")
        return
    await cb.answer()
    await save_and_next(bot, cb.message.chat.id, state, cb.from_user, data["step"], value)


@private.message(StateFilter(Form.phone), F.contact)
async def on_contact(message: Message, state: FSMContext, bot: Bot) -> None:
    phone = norm_phone(message.contact.phone_number) or message.contact.phone_number
    await save_and_next(bot, message.chat.id, state, message.from_user, "phone", phone)


@private.message(StateFilter(*STEP_STATES.values()), F.text, ~F.text.startswith("/"))
async def on_step_text(message: Message, state: FSMContext, bot: Bot) -> None:
    data = await state.get_data()
    step = data["step"]
    text = message.text.strip()
    if step == "phone" and text == BTN_BACK:
        await go_back(bot, message.chat.id, state, message.from_user, step)
        return
    if step == "phone" and text == BTN_SKIP:
        await save_and_next(bot, message.chat.id, state, message.from_user, step, None)
        return
    value = NORMALIZERS[step](text)
    if value is None:
        await message.answer(ERRORS[step])
        return
    await save_and_next(bot, message.chat.id, state, message.from_user, step, value)


@private.callback_query()
async def on_stale_callback(cb: CallbackQuery) -> None:
    await cb.answer("Эта кнопка уже неактуальна. Нажмите /start")


@private.message(STEP_FILTER)
async def on_step_other(message: Message) -> None:
    await message.answer("Пожалуйста, ответьте текстом или воспользуйтесь кнопками под вопросом.")


@private.message(StateFilter(Broadcast.confirm))
async def on_broadcast_confirm_other(message: Message) -> None:
    await message.answer("Воспользуйтесь кнопками под сообщением выше — «Разослать» или «Отмена».")


@private.message()
async def on_anything(message: Message) -> None:
    pending = active_application(message.from_user.id)
    if pending:
        await message.answer(pending_notice(pending))
        return
    await message.answer(
        "Чтобы оставить заявку, нажмите кнопку ниже 👇", reply_markup=apply_keyboard()
    )


async def setup_commands(bot: Bot) -> None:
    await bot.set_my_commands(
        [BotCommand(command="start", description="Оставить заявку")],
        scope=BotCommandScopeDefault(),
    )
    admin_commands = [
        BotCommand(command="start", description="Оставить заявку"),
        BotCommand(command="stats", description="Статистика"),
        BotCommand(command="export", description="Выгрузить заявки в Excel"),
        BotCommand(command="links", description="Ссылки для рекламы"),
        BotCommand(command="broadcast", description="Рассылка пользователям"),
    ]
    for admin_id in ADMIN_USER_IDS:
        with suppress(TelegramBadRequest):
            await bot.set_my_commands(admin_commands, scope=BotCommandScopeChat(chat_id=admin_id))


async def main() -> None:
    logging.basicConfig(level=logging.INFO)
    if not ADMIN_IDS:
        logging.warning("ADMIN_CHAT_IDS не задан: заявки будут только в базе %s", DB_PATH)
    if not ADMIN_USER_IDS:
        logging.warning("ADMIN_USER_IDS не задан: команда /stats будет недоступна никому")
    init_db()
    bot = Bot(BOT_TOKEN, default=bot_props)
    await setup_commands(bot)
    reminders = asyncio.create_task(reminder_loop(bot)) if REMIND_AFTER_MIN > 0 else None
    try:
        await dp.start_polling(bot)
    finally:
        if reminders:
            reminders.cancel()


if __name__ == "__main__":
    asyncio.run(main())
